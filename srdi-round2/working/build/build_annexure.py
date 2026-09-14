import sys, os, json
sys.path.insert(0, os.path.dirname(__file__))
from docx_lib import *
from content import *
from openpyxl import load_workbook
SP = "/tmp/claude-0/-home-user-Dataforge/8c3116aa-5914-54b9-9102-e947be44cbb4/scratchpad"
OUT = f"{SP}/out/Raunak_TeamWookies_NITWarangal_Annexure.docx"
wb = load_workbook(f"{SP}/out/recalc/H2SHIFT_model.xlsx", data_only=True)
srcs = json.load(open(f"{SP}/build/sources.json"))
d = new_doc(landscape=True, base_size=10)

def fmt(v, nd=1):
    if isinstance(v, (int, float)):
        if abs(v) >= 1000: return f"{v:,.0f}"
        return f"{v:,.{nd}f}"
    return "" if v is None else str(v)

# cover
para(d, "H2-SHIFT: Convert. Don't Replace.", size=22, bold=True, color=NAVY, space_after=2)
para(d, "Round 2 Annexure: supporting evidence, calculations and plans", size=14, bold=True, color=GREEN, space_after=8)
para(d, f"{TEAM['team']} | {TEAM['members']} | {TEAM['institute']}", size=11)
para(d, "SRDI Igniters Innovation Challenge 2026, Engineering Track, Option 2: Hydrogen Retrofit Solution. Research cut-off 14 September 2026.", size=10, color=GREY)
para(d, "This annexure supports the 10-slide main presentation and the workbook H2SHIFT_model.xlsx. All numbers here come from the recalculated workbook (base scenario) or from the source register. Evidence classes: A verified external evidence; B our calculated estimate; C explicit assumption or sensitivity range; D proposed target requiring validation. No customer interviews, letters of interest, test results, supplier quotations or partnerships are claimed; proposed partners are proposals.", size=10)
para(d, "Contents: A1 Claim audit | A2 Assumptions and source register | A3 Cost model | A4 Packaging, energy balance and performance | A5 Regulatory matrix and hazard analysis | A6 Market sizing | A7 Global benchmarks | A8 Pilot design and first 90 days | A9 Budget, responsibilities and funding | A10 Judge Q&A, validation log and change log.", size=10, color=GREY)

# A1
d.add_page_break(); d.add_heading("A1. Claim audit: Round 1 versus evidence", 1)
para(d, "Being shortlisted is evidence of competition progress, not proof of technical performance or commercial feasibility. Each Round 1 claim was checked against the research; the table gives the verdict and its Round 2 consequence.", size=9.5)
table(d, [["Original claim (Round 1)", "Evidence provided in Round 1", "Verification result", "Keep / refine / correct", "Round 2 implication"]] + [list(r) for r in CLAIM_AUDIT], widths=[5.2, 3.2, 8.6, 2.4, 6.6], size=7.5)

# A2
d.add_page_break(); d.add_heading("A2. Assumptions, evidence classes and source register", 1)
d.add_heading("A2.1 Model inputs (base, downside for the retrofit, upside)", 2)
ws = wb["Inputs"]; rows = [["Key", "Parameter", "Base", "Downside", "Upside", "Unit", "Class", "Sources", "Notes"]]
for r in range(6, ws.max_row + 1):
    k = ws.cell(row=r, column=1).value
    if not k: continue
    if ws.cell(row=r, column=6).value is None:
        rows.append([k, "", "", "", "", "", "", "", ""]); continue
    rows.append([k, ws.cell(row=r, column=2).value, fmt(ws.cell(row=r, column=3).value, 3), fmt(ws.cell(row=r, column=4).value, 3), fmt(ws.cell(row=r, column=5).value, 3), ws.cell(row=r, column=7).value, ws.cell(row=r, column=8).value, ws.cell(row=r, column=9).value, ws.cell(row=r, column=10).value])
table(d, rows, widths=[2.4, 7.0, 1.6, 1.6, 1.6, 1.8, 1.0, 2.6, 6.4], size=7)
d.add_heading("A2.2 Source register (309 entries; access date 14 September 2026)", 2)
para(d, "Prefix W1 engineering, W2 safety and regulation, W3 market, infrastructure and benchmarks, W4 economics and emissions. 'Excerpt' or 'summary' in Limitations means the original could not be opened through the build environment's network policy and the entry rests on indexed text of the page; re-verify before external use. Additional gap-closure sources read in full through the search index: PIB PRID 2006052 (14 Feb 2024), PIB PRID 2107795 (3 Mar 2025), Scribd copy of AIS-157 scope, Fuel Cells Works and RushLane on the Suzuki-AVL hydrogen Swift (May 2026), Luxfer hydrogen product page, Hexagon Purus Type 4 table (PDF read).", size=8.5)
rows = [["ID", "Claim supported", "Publisher", "Title", "URL", "Date", "Section", "Status", "Limitations"]]
for s_ in srcs:
    rows.append([s_.get("id", ""), s_.get("claim", ""), s_.get("publisher", ""), s_.get("title", ""), s_.get("url", ""), s_.get("date", ""), s_.get("section", ""), s_.get("status", ""), s_.get("limits", "")])
table(d, rows, widths=[1.4, 4.6, 2.6, 4.2, 6.2, 1.6, 1.8, 1.6, 2.6], size=6)

# A3
d.add_page_break(); d.add_heading("A3. Cost model: bill of materials, ownership cost, break-even and sensitivities", 1)
d.add_heading("A3.1 Retrofit bill of materials (pilot batch versus commercial volume)", 2)
ws = wb["BOM"]; rows = [["#", "Component or activity", "Qty", "Unit cost pilot (₹)", "Unit cost volume (₹)", "Pilot total (₹)", "Volume total (₹)", "Class", "Basis"]]
for r in range(5, 45):
    if ws.cell(row=r, column=2).value is None: continue
    rows.append([fmt(ws.cell(row=r, column=1).value, 0), ws.cell(row=r, column=2).value, fmt(ws.cell(row=r, column=3).value, 2), fmt(ws.cell(row=r, column=4).value, 0), fmt(ws.cell(row=r, column=5).value, 0), fmt(ws.cell(row=r, column=6).value, 0), fmt(ws.cell(row=r, column=7).value, 0), ws.cell(row=r, column=8).value or "", ws.cell(row=r, column=9).value or ""])
table(d, rows, widths=[0.7, 7.0, 1.0, 2.0, 2.0, 2.0, 2.0, 1.0, 8.0], size=7)
para(d, f"Installed conversion cost per vehicle including taxes: pilot ₹{B['pilot']:,.0f}; commercial volume ₹{B['volume']:,.0f}. Cylinder and valve are {B['cyl_share']*100:.0f} percent of parts cost at volume. One-time development and type-approval costs sit in the programme budget (A9) and are amortised at ₹8,000 per kit over 50,000 kits in the volume case.", size=9)
d.add_heading("A3.2 Ownership cost comparison (base scenario)", 2)
para(d, "Owner decision in 2028 with a 3-year-old Dzire Tour S CNG; horizon 5 years; 55,000 km per year (167 km x 330 days); nominal discount rate 12 percent; fuel escalation 3 percent, electricity 2 percent, delivered hydrogen minus 10 percent per year from ₹800/kg. Option A keeps the car (residual value at year 5), option B converts it (kit at t0, residual at year 5 reduced by 10 percent), option C sells it at t0 (₹3.5 lakh) and buys a compact BEV (₹10.5 lakh, residual ₹4.2 lakh at year 5). In-shift refuelling or charging time is costed in all three options at ₹150 per hour; overnight charging is free idle time. No depreciation is charged in addition to capital; no battery replacement is assumed (8-year, 160,000 km warranties).", size=9)
T = O["TCO_rows"]
rows = [["Total cash cost by year (₹)", "t0 (2028)", "2029", "2030", "2031", "2032", "2033"]]
for lab, key in [("A: continue on CNG", "A_total"), ("B: hydrogen retrofit", "B_total"), ("C: sell and buy BEV", "C_total")]:
    rows.append([lab] + [fmt(v, 0) for v in T[key]])
rows.append(["Delivered hydrogen price path (₹/kg)", ""] + [fmt(v, 0) for v in T["h2price"]])
table(d, rows, widths=[6, 3, 3, 3, 3, 3, 3], size=8, bold_first_col=True)
rows = [["Result (base scenario)", "Value"],
 ["NPV of costs: A / B / C (₹ lakh)", f"{S['npvA']/1e5:.1f} / {S['npvB']/1e5:.1f} / {S['npvC']/1e5:.1f}"],
 ["Levelised cost per km: A / B / C (₹/km)", f"{S['lcA']:.2f} / {S['lcB']:.2f} / {S['lcC']:.2f}"],
 ["Year-1 fuel or energy cost per km: CNG mode / hydrogen mode / petrol mode / BEV (₹/km)", f"{S['fkm_cng']:.2f} / {S['fkm_h2']:.2f} / {S['fkm_petrol']:.2f} / {S['fkm_bev']:.2f}"],
 ["Year-1 blended fuel cost per km: A / B (₹/km)", f"{S['fkm_A']:.2f} / {S['fkm_B']:.2f}"],
 ["Year-1 non-fuel operating cost per km: A / B / C (₹/km)", f"{S['okm_A']:.2f} / {S['okm_B']:.2f} / {S['okm_C']:.2f}"],
 ["Extra cost of B versus A; B versus C; C versus A over 5 years (₹ lakh)", f"{S['dBA']/1e5:.1f}; {S['dBC']/1e5:.1f}; {S['dCA']/1e5:.1f}"],
 ["Fuel-parity hydrogen price versus CNG; versus petrol (₹/kg)", f"{S['parity_cng']:.0f}; {S['parity_petrol']:.0f}"],
 ["Full-TCO break-even year-1 hydrogen price versus A; versus C (₹/kg)", f"{S['be_A']:.0f} (negative: no positive price breaks even); {S['be_C']:.0f}"],
 ["Maximum conversion cost for parity with A; with C at the base hydrogen price (₹)", f"{S['maxkit_A']:,.0f}; {S['maxkit_C']:,.0f}"],
 ["Simple payback of the conversion versus A", str(S['payback'])],
 ["Minimum annual km for payback within the horizon", str(S['minkm'])],
 ["Decomposition check (should be zero)", f"{S['chk']:.2f} and {S['chk2']:.2f}"],
]
table(d, rows, widths=[13, 12], size=8.5, bold_first_col=True)
d.add_heading("A3.3 Sensitivity grids (₹ lakh extra cost over 5 years; positive means the retrofit costs more)", 2)
SN = O["SENS"]
def grid(title, g, colfmt):
    para(d, title, size=9, bold=True, color=NAVY, space_after=2)
    rows = [[str(g[0][0])] + [colfmt(c) for c in g[0][1:]]]
    for r in g[1:]:
        rows.append([f"₹{r[0]:.0f}/kg"] + [f"{v:+.1f}" for v in r[1:]])
    table(d, rows, size=8, bold_first_col=True)
grid("Grid 1: retrofit minus continue-on-CNG by year-1 hydrogen price (rows) and annual km (columns)", SN["grid1"], lambda c: f"{c:,.0f} km")
grid("Grid 2: retrofit minus sell-and-buy-BEV by hydrogen price (rows) and BEV price in ₹ lakh (columns)", SN["grid2"], lambda c: f"BEV ₹{c:.0f} lakh")
grid("Grid 3: retrofit minus continue-on-CNG by hydrogen price (rows) and installed kit cost in ₹ lakh (columns)", SN["grid3"], lambda c: f"kit ₹{c:.1f} lakh")
rows = [["One-at-a-time driver", "Low value", "High value", "Extra cost at low (₹ lakh)", "Extra cost at high (₹ lakh)"]] + [[r[0], fmt(r[1], 2), fmt(r[2], 2), f"{r[3]:+.1f}", f"{r[4]:+.1f}"] for r in SN["oat"]]
table(d, rows, widths=[8, 3, 3, 4, 4], size=8, bold_first_col=True)
para(d, "Reading: the extra cost of the retrofit stays positive across every grid cell; the smallest values occur at low hydrogen prices, low mileage and low kit cost. The BEV comparison is closest at high BEV prices and low hydrogen prices.", size=9)

# A4
d.add_page_break(); d.add_heading("A4. Packaging, energy balance and performance calculations", 1)
ws = wb["Energy_Range"]; rows = [["Quantity", "Value", "Unit", "Method or note"]]
for r in range(5, 35):
    q = ws.cell(row=r, column=1).value
    if q: rows.append([q, fmt(ws.cell(row=r, column=2).value, 3), ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value or ""])
table(d, rows, widths=[9, 2.6, 2.4, 12], size=8)
para(d, "Packaging basis (concept, to be verified by mock-up): the Dzire Tour S carries a 55 L Type I steel CNG cylinder (OD 316 to 325 mm, about 0.85 to 1.0 m long, 48 to 69 kg) transversely in the boot. The replacement is a 60 L composite cylinder of the Luxfer G-Stor Pro V060H35 class (Type III, 350 bar, OD 400 mm, 767 mm long, 36 kg, 1.45 kg hydrogen; Type IV equivalents from Quantum or NPROXX exist in the 26 to 80 L range but no accessible datasheet was found). The larger diameter consumes an estimated further 40 to 60 L of luggage space; seating is unchanged; ground clearance of 163 mm rules out under-floor mounting. Mounting frame designed to the UN R110 M1 loads of 20 g longitudinal and 8 g lateral (hydrogen-specific installation clauses of AIS-195 to be applied once accessible). Non-cooled 350 bar fills are assumed to reach 85 percent state of charge in hot ambient conditions; SAE J2601 light-duty tables start at 49.7 L storage volume, so the fill protocol must be agreed with the station.", size=9)
para(d, "Consumption cross-check: the fleet CNG car at 25 km/kg uses 1.92 MJ/km; the hydrogen car at 1.85 kg/100 km uses 2.22 MJ/km, an efficiency ratio of 0.87 (lean port-injection retrofit calibration below CNG stoichiometric efficiency). Comparable vehicles: Ford Focus H2RV 1.4 kg/100 km (supercharged hybrid, claimed), Mazda RX-8 Hydrogen RE 2.4 kg/100 km (rotary), BMW Hydrogen 7 3.7 kg/100 km (V12, FTP-75), Toyota hydrogen HiAce about 3.1 kg/100 km (3.4 L V6 van, derived). Power: port-injected hydrogen displaces about 30 percent of intake air (about 20 percent power loss at stoichiometric) and lean operation at lambda 2 costs a further 35 to 40 percent; the model uses 52 percent of petrol power (42 PS) with lambda 1.6 to 2.2.", size=9)

# A5
d.add_page_break(); d.add_heading("A5. Regulatory compliance matrix and hazard analysis", 1)
d.add_heading("A5.1 Compliance matrix", 2)
table(d, [["Requirement", "Applicable technology / category", "Current source", "Required evidence or test", "Responsible party", "Status / uncertainty"]] + [list(r) for r in COMPLIANCE], widths=[4.2, 3.4, 6.2, 4.6, 3.4, 4.6], size=7.5)
d.add_heading("A5.2 Hazard analysis (prevention, detection and response, validation, residual risk)", 2)
table(d, [["Hazard", "Prevention", "Detection and response", "Validation", "Residual risk"]] + [list(r) for r in HAZARDS], widths=[3.6, 6.6, 5.2, 5.6, 5.4], size=7.5)
para(d, "Hydrogen properties used: flammability 4.0 to 75.0 percent in air; detonation 18.3 to 59 percent; minimum ignition energy 0.017 mJ; autoignition 585 C; diffusivity about 3.8 times natural gas; nearly invisible flame (H2Tools, US DOE). Prototype work and high-pressure testing are confined to qualified facilities (test cells at HORIBA Chakan, ARAI, IOCL R&D; PESO-approved test stations). This annexure is a development and validation plan, not conversion instructions.", size=9)

# A6
d.add_page_break(); d.add_heading("A6. Market sizing method and bottom-up funnel", 1)
rows = [["Step", "Vehicles (base)", "Share applied", "Basis"]]
shares = ["", f"{IN['cng_pv_share']['active']*100:.0f} percent", f"{IN['eligible_share']['active']*100:.0f} percent", f"{IN['fleet_share']['active']*100:.0f} percent", f"{IN['hub_share']['active']*100:.0f} percent", f"{IN['adopter_share']['active']*100:.0f} percent", "", "", ""]
basis = ["PNGRB citing Vahan (Mar 2025), class A", "Estimate from cumulative factory-CNG car sales (Maruti about 2.6 million FY20-FY26 plus other OEMs) and conversions; Round 1 said 65 percent (unverified)", "Factory-fitted, boot cylinder, under 6 years old at 2028 (FY23-FY26 sales)", "Delhi alone registers 60,000 to 85,000 taxis; no national split published; assumption", "MNRE corridors: Delhi-NCR, Ahmedabad-Vadodara-Surat, Pune-Mumbai, Kochi and others", "Hypothesis to be validated in Phase 0", "8 hub stations x 174 cars per station (500 kg/day at 70 percent utilisation, 2.0 kg per car per day)", "Minimum of demand and supply", "For comparison"]
for (lab, val), sh, ba in zip(M["funnel"], shares, basis):
    rows.append([lab, fmt(val, 0), sh, ba])
table(d, rows, widths=[8, 3, 2.5, 12.5], size=8, bold_first_col=True)
para(d, "Downside and upside funnels (workbook scenario selector) give roughly 700 to 2,600 cars by 2030; in every case supply (stations) rather than demand binds. Private owners are excluded from the first segment because they lack depot fuelling and demand aggregation; they could follow only after hub stations exist.", size=9)

# A7
d.add_page_break(); d.add_heading("A7. Global benchmark evidence: progress and setbacks", 1)
table(d, [["Programme", "Technology and class", "Status (date)", "Scale", "Fuel-supply model", "Commercial limitation", "What transfers to an Indian retrofit"]] + [list(r) for r in BENCHMARKS], widths=[3.6, 3.6, 4.6, 2.8, 3.0, 3.6, 5.2], size=7.5)

# A8
d.add_page_break(); d.add_heading("A8. Pilot design, instrumentation and first 90 days", 1)
table(d, [["Element", "Design"]] + [list(r) for r in PILOT_DESIGN], widths=[5, 21], size=8.5, bold_first_col=True)
d.add_heading("A8.1 Success measures for gate 4 (September 2030)", 2)
bullets(d, ["Consumption at or below 2.2 kg/100 km in fleet use; range per fill at or above 55 km; hydrogen share of km at or above 60 percent.",
            "Availability at or above 95 percent; conversion-related downtime below 2 percent of operating days.",
            "Zero reportable hydrogen incidents; all alarms investigated; no unplanned venting.",
            "Tailpipe NOx at or below 60 mg/km on the certification cycle and on-road spot checks within 1.5 times the limit.",
            "Measured service and maintenance cost within ₹1.2/km; refuelling time under 6 minutes per fill; luggage and passenger acceptance from fleet managers.",
            "Delivered hydrogen price at or below ₹250/kg with a documented path to ₹190/kg; station utilisation trend above 50 percent."], size=9)
d.add_heading("A8.2 First 90 days (from programme start, October 2026)", 2)
table(d, [["Window", "Actions"]] + [list(r) for r in NINETY_DAYS], widths=[3.5, 22.5], size=8.5, bold_first_col=True)

# A9
d.add_page_break(); d.add_heading("A9. Phase budgets, responsibilities and funding", 1)
rows = [["Phase / line item", "Quantity", "Unit cost (₹)", "Total (₹)"]]
for r in PB["lines"]:
    ph, name, qty, uc, tot_ = r
    if ph and not name: rows.append([ph, "", "", ""])
    elif name: rows.append([name, fmt(qty, 1) if qty is not None else "", fmt(uc, 0) if uc is not None else "", fmt(tot_, 0) if tot_ is not None else ""])
table(d, rows, widths=[15, 3, 4, 4], size=8)
para(d, f"Programme total to September 2030 including 15 percent contingency: ₹{tot_cr:.1f} crore (phases ₹{ph_cr[0]:.1f}, ₹{ph_cr[1]:.1f}, ₹{ph_cr[2]:.1f}, ₹{ph_cr[3]:.1f} crore). Proposed funding split: Suzuki R&D Center India and Maruti Suzuki 40 percent; public support 35 percent (MNRE transport pilot scheme, which names buses, trucks and 4-wheelers, or the R&D scheme with up to 80 percent support for industry at TRL 4 or higher; Gujarat's 30 percent station subsidy for the second hub); oil marketing company partner 15 percent; fleet partner 10 percent. For calibration, MNRE's heavy-vehicle pilots average ₹5.6 to 5.9 crore of support per vehicle including stations. All partners are proposed, not confirmed.", size=9)
rows = [["Responsibility", "Proposed holder", "Rationale"],
 ["Engineering lead, calibration, prototypes", "Suzuki R&D Center India with Maruti Suzuki engineering", "CNG engine know-how; Suzuki hydrogen combustion prototypes"],
 ["Kit integration, warranty, installer training", "Kit integrator (aftermarket gaseous-fuel specialist) under SRDI oversight", "Existing CNG kit approval experience"],
 ["Type approval and testing", "ARAI or ICAT (Rule 126 agencies); PESO for cylinders and station", "Statutory"],
 ["Dispensing point ownership and operation", "Oil marketing company (proposed IOCL, Faridabad)", "Existing hydrogen dispensing experience; corridor awardee"],
 ["Hydrogen supply", "SIGHT-awarded green hydrogen producer by tube trailer", "Certified carbon intensity"],
 ["Vehicles and operations", "Fleet partner (PSU or corporate staff transport, airport contract)", "Return-to-base duty"],
 ["Independent evaluation", "IIT Delhi or CEEW (proposed)", "Credibility of results"],
 ["Development risk", "Programme (consortium)", "Gated spend"], ["Utilisation risk", "Station partner with take-or-pay from the fleet", "Aligns incentives"], ["Residual-value risk", "Fleet, backed by programme buy-back", "Removes owner downside in the pilot"]]
table(d, rows, widths=[7, 9, 10], size=8.5, bold_first_col=True)

# A10
d.add_page_break(); d.add_heading("A10. Judge Q&A, validation log and change log", 1)
d.add_heading("A10.1 Judge Q&A (20 questions)", 2)
table(d, [["Question", "Answer"]] + [list(r) for r in QA], widths=[7, 19], size=8, bold_first_col=True)
d.add_heading("A10.2 Validation log", 2)
table(d, [["Item", "What was done"]] + [list(r) for r in VALIDATION_LOG], widths=[4, 22], size=8.5, bold_first_col=True)
d.add_heading("A10.3 Change log: Round 1 to Round 2", 2)
table(d, [["Element", "Round 1", "Round 2"]] + [list(r) for r in CHANGE_LOG], widths=[4, 10, 12], size=8.5, bold_first_col=True)
para(d, TEAM["leader_note"], size=9, italic=True, color=GREY)
d.save(OUT); print("saved", OUT)
