"""v2: dense consulting-style redesign. Slide 1 exact; slides 2-10 exhibit-led. Numbers from model_outputs_base.json and chart_data_v2.json."""
import json, sys, os, math
sys.path.insert(0, os.path.dirname(__file__))
from deck_lib import *
SP = "/tmp/claude-0/-home-user-Dataforge/8c3116aa-5914-54b9-9102-e947be44cbb4/scratchpad"
O = json.load(open(f"{SP}/out/model_outputs_base.json"))
CD = json.load(open(f"{SP}/out/chart_data_v2.json"))
S, E, H, M, EM, B, PB, IN = O["S"], O["ER"], O["H2S"], O["MK"], O["EM"], O["BOM"], O["PB"], O["INPUTS"]
bd = H["build"]
OUT = sys.argv[1] if len(sys.argv) > 1 else f"{SP}/v2/Raunak_TeamWookies_NITWarangal.pptx"
AMBER = RGBColor(0xE0, 0x9A, 0x1E); REDC = RGBColor(0xB0, 0x3A, 0x2E); GREYC = RGBColor(0x8A, 0x8D, 0x8F); BLUEC = TAGBLUE
STATUS = {"Established": GREEN, "Unresolved": AMBER, "Gap": REDC, "Target": AMBER, "Not met": AMBER, "Hypothesis": GREYC, "Prototype": GREYC, "Pilot": BLUEC, "Commercial": GREEN, "Withdrawn": REDC, "Shrinking": AMBER, "Approved": GREEN}
def L(x, d=1): return f"{x/1e5:.{d}f}"
def n0(x): return f"{x:,.0f}"
tot_cr = PB["total"] / 1e7; ph_cr = [x / 1e7 for x in PB["phases"]]

def header2(slide, headline, subline=None):
    n = len(headline)
    size = 28 if n <= 120 else 26 if n <= 150 else 24
    textbox(slide, MARGIN, Inches(0.32), Inches(16.0), Inches(1.15), headline, size=size, color=NAVY, bold=True, line_spacing=1.0, space_after=0)
    if subline:
        textbox(slide, MARGIN, Inches(1.38), Inches(16.3), Inches(0.5), subline, size=14.5, color=DARKGREY, line_spacing=1.0, space_after=0)
    tb = slide.shapes.add_textbox(Inches(16.8), Inches(0.32), Inches(2.8), Inches(0.9)); tf = tb.text_frame; tf.margin_left = tf.margin_right = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    for t, c in (("H", SUZUKIBLUE), ("₂", SUZUKIBLUE), ("-SHIFT", DEEPGREEN)):
        r = p.add_run(); r.text = t; r.font.size = Pt(26); r.font.bold = True; r.font.italic = True; r.font.name = FONT; r.font.color.rgb = c
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.RIGHT
    r = p2.add_run(); r.text = "Convert. Don't Replace."; r.font.size = Pt(11.5); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = NAVY
    line(slide, MARGIN, Inches(1.9), SLIDE_W - MARGIN, Inches(1.9), color=GREEN, width=1.5)
    return Inches(2.05)

def footer2(slide, text, no):
    label(slide, MARGIN, Inches(10.62), Inches(17.6), Inches(0.5), text, size=10.5, color=MIDGREY)
    label(slide, Inches(18.9), Inches(10.78), Inches(0.6), Inches(0.3), str(no), size=10.5, color=MIDGREY, align=PP_ALIGN.RIGHT)

def kv_table(slide, x, y, w, rows, kw=1.35, size=12, row_h=Inches(0.5), key_color=DEEPNAVY):
    """Two-column key/value list with light rules."""
    cy = y
    for k, v in rows:
        label(slide, x, cy, Inches(kw), row_h, k, size=size, color=key_color, bold=True, anchor=MSO_ANCHOR.TOP)
        label(slide, x + Inches(kw), cy, w - Inches(kw), row_h, v, size=size, color=DARKGREY, anchor=MSO_ANCHOR.TOP)
        cy += row_h
        line(slide, x, cy - Inches(0.04), x + w, cy - Inches(0.04), color=GREY2, width=0.5)
    return cy

def status_chip(slide, x, y, text, w=Inches(1.25), h=Inches(0.3), size=10.5):
    key = text.split(":")[0].split(" ")[0]
    fill = STATUS.get(text, STATUS.get(key, GREYC))
    chip(slide, x, y, w, h, text, fill, size=size)

def table2(slide, x, y, w, h, data, col_widths=None, font_size=12, header_size=12, highlight_col=None, row_heights=None, align_center_from=99, bold_rows=(), first_col_bold=True):
    return table(slide, x, y, w, h, data, col_widths=col_widths, font_size=font_size, header_size=header_size, highlight_col=highlight_col,
                 row_heights=row_heights, align_center_from=align_center_from, bold_rows=bold_rows, first_col_bold=first_col_bold)

prs = new_presentation()
add_title_slide(prs, f"{SP}/assets/r1page-000.png")

# =============================================================== SLIDE 2
s = blank_slide(prs)
y0 = header2(s, "Recommendation: validate before selling. Run a 30-car hydrogen retrofit pilot in a captive Faridabad fleet and scale only if delivered hydrogen falls below about ₹190 per kg",
             "The brief asks whether retrofits can be a safe, affordable, viable complement to BEVs by 2030. Evidence answer: safe with engineering, not affordable at unsubsidised hydrogen prices, viable only as a station-limited niche.")
# ---- left: funnel
lx, lw = MARGIN, Inches(7.3)
cy = section_title(s, lx, y0, lw, "Stations, not demand, set the niche: bottom-up funnel to 2030 (model)")
fun = M["funnel"][:8]
labels_f = ["CNG vehicles on road, Mar 2025", "Passenger cars (estimate, 55 percent)", "Factory-fitted, boot bay, under 6 yr in 2028", "In commercial or institutional fleets", "In cities with a hydrogen hub by 2030", "Early adopters, return-to-base (hypothesis)", "Station-limited capacity: 8 hubs x 174 cars", "Realistic converted fleet, 2030"]
maxlog = math.log10(fun[0][1]); minlog = math.log10(1000)
rh = Inches(0.55); barx = lx + Inches(3.55); barw = lw - Inches(3.55)
for i, ((lab, val), name) in enumerate(zip(fun, labels_f)):
    ry = cy + i * (rh + Inches(0.09))
    label(s, lx, ry, Inches(3.5), rh, name, size=12, color=DARKGREY if i < 7 else DEEPNAVY, bold=(i == 7), anchor=MSO_ANCHOR.MIDDLE)
    frac = (math.log10(max(val, 1000)) - minlog) / (maxlog - minlog)
    fill = GREEN if i == 7 else (CHARTBLUE if i < 5 else TAGBLUE)
    hbar_row(s, barx, ry + Inches(0.08), barw, rh - Inches(0.16), frac, fill, text_right=n0(val), size=12)
fy = cy + 8 * (rh + Inches(0.09)) + Inches(0.05)
label(s, lx, fy, lw, Inches(0.55), "Bar length on a log10 scale (1,000 to 8.2 million). One 500 kg/day hub at 70 percent utilisation serves 174 cars at 2.0 kg per car per day; demand (9,000) exceeds station capacity (1,394) sevenfold. Round 1 said 1,00,000 cars.", size=11, color=MIDGREY)
# answer to the brief: three cards
ay = fy + Inches(0.72); cw = (lw - Inches(0.3)) / 3
cards = [("SAFE: yes, with engineering", GREEN, "Composite cylinder, thermal relief, twin sensors, lean burn with NOx control. No approval route exists for a hydrogen retrofit today: gate 1."),
         ("AFFORDABLE: not by 2030", REDC, f"₹{S['lcB']:.1f}/km converted vs ₹{S['lcA']:.1f}/km on CNG at ₹{IN['h2_price_y1']['active']:.0f}/kg (2028). CNG fuel parity needs ₹{S['parity_cng']:.0f}/kg; cheapest 2030 forecast about ₹360/kg."),
         ("COMPLEMENT: niche only", AMBER, f"A compact BEV costs ₹{S['lcC']:.1f}/km including purchase. Hydrogen retrofits fit captive fleets beside hydrogen hubs, about 1,400 cars by 2030.")]
for i, (t, col, body) in enumerate(cards):
    cx = lx + i * (cw + Inches(0.15))
    rect(s, cx, ay, cw, Inches(2.15), RGBColor(0xF6, 0xF7, 0xF8)); rect(s, cx, ay, cw, Inches(0.36), col)
    label(s, cx + Inches(0.08), ay, cw - Inches(0.1), Inches(0.36), t, size=11.5, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    label(s, cx + Inches(0.08), ay + Inches(0.44), cw - Inches(0.16), Inches(1.7), body, size=12, color=DARKGREY)
# ---- middle: customer card
mx, mw = lx + lw + Inches(0.3), Inches(5.55)
cy = section_title(s, mx, y0, mw, "Initial customer and use case")
kv_rows = [("Donor", "Maruti Suzuki Dzire Tour S CNG, 3 to 5 years old, factory bi-fuel; petrol stays as automatic backup"),
           ("Fleet", "Captive, return-to-base institutional and contract fleets: PSU and corporate staff transport, airport contracts"),
           ("Duty", "150 to 210 km per day, 330 days; two depot fills per day give about 120 km on hydrogen"),
           ("Location", "Faridabad, Delhi-NCR: IOCL R&D hydrogen dispensing point and the MNRE Sahibabad-Faridabad-Delhi corridor (sanctioned Mar 2025)"),
           ("Excluded", "Delhi app-cabs: aggregator scheme counts only EVs; CAQM allows only CNG or EV new aggregator cars from Jan 2026"),
           ("Second hub", "Ahmedabad-Vadodara-Surat corridor; Gujarat funds 30 percent of the first 20 hydrogen stations")]
cy = kv_table(s, mx, cy, mw, kv_rows, kw=1.15, size=12, row_h=Inches(0.78))
cy = section_title(s, mx, cy + Inches(0.12), mw, "Decision metrics (model, 2028 decision, 5 years)")
kw_, kh_ = (mw - Inches(0.15)) / 2, Inches(1.22)
kpis = [(f"₹{S['lcB']:.1f} vs {S['lcA']:.1f} vs {S['lcC']:.1f}", "per km: retrofit vs stay on CNG vs BEV (levelised, all-in)"),
        (f"₹{S['parity_cng']:.0f}/kg", "delivered hydrogen price for fuel parity with CNG"),
        (f"{E['range']:.0f} km", f"per hydrogen fill ({E['usable']:.2f} kg usable); {E['petrolrange']:.0f} km petrol backup"),
        (f"~{round(M['fleet2030'], -2):,.0f} cars", "conditional 2030 fleet, station-limited (Round 1: 1,00,000)")]
for i, (v, c) in enumerate(kpis):
    kpi(s, mx + (i % 2) * (kw_ + Inches(0.15)), cy + (i // 2) * (kh_ + Inches(0.12)), kw_, kh_, v, c, value_size=20, cap_size=11)
# ---- right: conditions
rx, rw = mx + mw + Inches(0.3), SLIDE_W - MARGIN - (mx + mw + Inches(0.3))
cy = section_title(s, rx, y0, rw, "Conditions for success and where they stand today")
cond = [("Retrofit approval route defined (gate 1, Mar 2027)", "None exists; AIS-195 covers new vehicles only", "Gap"),
        (f"Delivered hydrogen at or below ₹250/kg by 2030 with a path to ₹{S['parity_cng']:.0f}", f"Model hub ₹{bd['delivered'][1]:,.0f}; CEEW central ₹360", "Not met"),
        ("Certified kit at or below ₹2.5 lakh installed", f"₹{L(B['volume'])} lakh at 10,000/yr; ₹{L(B['pilot'])} lakh pilot", "Not met"),
        ("NOx at or below 60 mg/km on the BS-VI cycle", "Lean burn plus NOx trap provision; not yet tested", "Target"),
        ("Three fleets willing to accept 48 percent power loss and a smaller boot", "No interviews yet; Phase 0 buys 30", "Hypothesis")]
for i, (c1, c2, st) in enumerate(cond):
    ry = cy + i * Inches(0.92)
    label(s, rx, ry, rw - Inches(1.4), Inches(0.45), c1, size=12, color=DEEPNAVY, bold=True)
    label(s, rx, ry + Inches(0.45), rw - Inches(1.4), Inches(0.42), c2, size=11.5, color=DARKGREY)
    status_chip(s, rx + rw - Inches(1.3), ry + Inches(0.16), st, w=Inches(1.3))
    line(s, rx, ry + Inches(0.88), rx + rw, ry + Inches(0.88), color=GREY2, width=0.5)
cy = cy + 5 * Inches(0.92) + Inches(0.12)
cy = section_title(s, rx, cy, rw, "What changed since Round 1")
chg = [("Fallback fuel", "petrol already fitted, not a second CNG cylinder"), ("Price basis", "pump price (₹584 to 1,433/kg), not production cost (₹160 to 200)"), ("Rollout", "3 prototypes, 30 pilot cars, about 1,400 by 2030, not 1,00,000"), ("Stop rules", "four gates with proceed, redesign or stop criteria")]
kv_table(s, rx, cy, rw, chg, kw=1.25, size=11.5, row_h=Inches(0.56))
footer2(s, "Sources: PNGRB citing Vahan, 81.95 lakh CNG vehicles [W3-S2]; Delhi aggregator scheme [W3-S17]; CAQM direction [W3-S19]; IOCL Faridabad [W3-S46]; MNRE corridors, PIB 3 Mar 2025 [W3-S36]; Gujarat policy [W4-S76]; CEEW 2025 [W4-S22]; workbook sheets Market, TCO, Energy_Range.", 2)

# =============================================================== SLIDE 3
s = blank_slide(prs)
y0 = header2(s, "Technology choice: H2-ICE is the only route that is a retrofit in cost and scope, but its fuel economics are the weakest; a fuel-cell conversion costs as much as a new BEV",
             "Matched criteria for the same donor and duty: Dzire Tour S CNG, 55,000 km per year, owner decision in 2028. Model outputs unless a source ID is given; estimates are labelled.")
tw = Inches(12.85)
rows = [["Criterion", "H2-ICE retrofit (selected)", "Fuel-cell conversion", "New compact BEV, donor sold", "Continue on CNG"],
 ["Scope of change", "Cylinder, lines, injectors, ECU, sensors; engine and petrol system retained", "Whole powertrain: stack, battery, motor, inverter, cooling, tank, controls", "New vehicle; donor sold for about ₹3.5 lakh", "None"],
 ["Cost to the owner, 2028", f"₹{L(B['volume'])} lakh installed at 10,000/yr (₹{L(B['pilot'])} lakh pilot)", "₹15 to 25 lakh pilot; ₹8 to 12 lakh at volume (estimate, W4)", f"₹{L(IN['bev_price']['active'])} lakh (Tiago EV to Punch EV class) [W3-S27]", "₹0"],
 ["Energy cost per km, 2028", f"₹{S['fkm_B']:.1f} blended; hydrogen km ₹{S['fkm_h2']:.1f} at ₹{IN['h2_price_y1']['active']:.0f}/kg", "About half of H2-ICE per hydrogen km (stack efficiency)", f"₹{S['fkm_bev']:.1f} at ₹6 depot, ₹20 public per kWh", f"₹{S['fkm_A']:.1f} at ₹87/kg"],
 ["All-in cost per km, 5 years", f"₹{S['lcB']:.1f}", "Not modelled (capital above BEV)", f"₹{S['lcC']:.1f}", f"₹{S['lcA']:.1f}"],
 ["Range per fill or charge", f"{E['range']:.0f} km on hydrogen plus {E['petrolrange']:.0f} km petrol", "About 2x H2-ICE on the same tank", "187 to 260 km, real-world tests [W3-S30]", "About 250 km at 25 km/kg"],
 ["Refuel or recharge time", "3 to 5 min, 350 bar non-cooled fill", "3 to 5 min", "About 1 h DC 10 to 80 percent; 6 to 8 h AC", "3 to 5 min plus queue"],
 ["Peak power vs petrol", f"About {IN['h2_power_factor']['active']*100:.0f} percent, lean burn (city duty)", "Restored by electric drive", "Higher", "About 85 percent"],
 ["Tailpipe", "Zero CO2 on hydrogen km; NOx must be controlled", "Water vapour only", "Zero", "About 105 g CO2/km"],
 ["Approval route today", "None for retrofits; AIS-195 covers new H2-ICE vehicles", "None; AIS-157 scope excludes retrofitted FCEVs", "Established", "Established"],
 ["Supplier readiness", "Injectors, valves, regulators production-intent; car-size cylinders scarce", "Imported stacks; no light-duty supply chain in India", "Mature", "Mature"],
]
table2(s, MARGIN, y0, tw, Inches(6.3), rows, col_widths=[1.45, 2.35, 2.2, 2.1, 1.5], font_size=11.5, header_size=12, highlight_col=1)
# right: status ladder
rx = MARGIN + tw + Inches(0.25); rw = SLIDE_W - MARGIN - rx
cy = section_title(s, rx, y0, rw, "Global deployments: what is real, and the lesson")
ladder = [("Toyota Mirai (FCEV car)", "Commercial", "US sales 210 in 2025; fuel credits needed [W3-S51]"),
          ("Hyundai Nexo gen 2 (FCEV)", "Commercial", "Korea only at scale, subsidy-dependent [W3-S52]"),
          ("Toyota hydrogen HiAce (H2-ICE van)", "Pilot", "Depot fuelled; under 200 km; SCR for NOx [W3-S55]"),
          ("Suzuki-AVL hydrogen Swift (H2-ICE car)", "Prototype", "1.4 L DI, 100 kW, Vienna May 2026; not a retrofit"),
          ("KEYOU truck conversions (H2-ICE)", "Pilot", "Only commercial retrofit model; Daimler 2027 [W3-S61]"),
          ("JCB hydrogen engine (off-highway)", "Approved", "EU type approval May 2025 [W3-S60]"),
          ("Stellantis hydrogen vans (FCEV)", "Withdrawn", "Cancelled Jul 2025: infrastructure, capital [W3-S64]"),
          ("Shell CA, H2 Mobility DE car stations", "Withdrawn", "43 car stations closed 2024 to 2025 [W3-S65, S66]")]
for i, (p_, st, les) in enumerate(ladder):
    ry = cy + i * Inches(0.7)
    label(s, rx, ry, rw - Inches(1.3), Inches(0.34), p_, size=11.5, color=DEEPNAVY, bold=True)
    label(s, rx, ry + Inches(0.32), rw - Inches(1.3), Inches(0.36), les, size=10.5, color=DARKGREY)
    status_chip(s, rx + rw - Inches(1.2), ry + Inches(0.08), st, w=Inches(1.2), size=10)
    line(s, rx, ry + Inches(0.67), rx + rw, ry + Inches(0.67), color=GREY2, width=0.5)
label(s, rx, cy + 8 * Inches(0.7) + Inches(0.05), rw, Inches(0.5), "Status: Prototype = demonstrator; Pilot = fleet trial with data; Commercial = product on sale; Withdrawn = programme or stations closed.", size=10, color=MIDGREY)
# bottom strip: fit
by = y0 + Inches(6.5)
cy = section_title(s, MARGIN, by, SLIDE_W - 2 * MARGIN, "Why the selected route fits the donor, the customer and the question")
cw = (SLIDE_W - 2 * MARGIN - Inches(0.4)) / 3
fit = [("Donor fit", "1.2 L four-cylinder SI engine with CNG-hardened seats, a 55 L cylinder bay and a 37 L petrol tank: only the fuel system is new; engine, gearbox, body and petrol backup are retained."),
       ("Customer fit", f"Return-to-base duty with two depot fills makes a {E['range']:.0f} km hydrogen range workable for {E['maxshare']*100:.0f} percent of a 167 km day; petrol covers the rest, so no car is stranded by a missing station."),
       ("Economic fit", f"The only route whose cost is a kit (₹{L(B['volume'])} lakh) rather than a vehicle (₹8 to 25 lakh). Its case rests on one variable, the delivered hydrogen price, which the pilot is built to measure.")]
for i, (t, body) in enumerate(fit):
    cx = MARGIN + i * (cw + Inches(0.2))
    rect(s, cx, cy, Inches(0.07), Inches(1.45), GREEN)
    label(s, cx + Inches(0.18), cy, cw - Inches(0.2), Inches(0.32), t, size=12.5, color=DEEPNAVY, bold=True)
    label(s, cx + Inches(0.18), cy + Inches(0.32), cw - Inches(0.2), Inches(1.15), body, size=11.5, color=DARKGREY)
footer2(s, "Sources: Donor_Tech and TCO sheets; fuel-cell conversion cost is an estimate from stack and battery cost literature [W4]; BEV data [W3-S27, S30]; AIS-157 scope [W2-S11, Scribd copy]; programmes [W3-S51 to S66]; Suzuki-AVL Swift (Fuel Cells Works, RushLane, May 2026).", 3)

# =============================================================== SLIDE 4
s = blank_slide(prs)
y0 = header2(s, "Retrofit engineering: storage, fuel delivery, controls and safety are new; the engine, petrol system and body are retained. One fill gives about 61 km; power falls by about half",
             "Donor: Maruti Suzuki Dzire Tour S CNG, 1.2 L K-series, 60 kW on petrol, 55 L steel CNG cylinder in the boot, 37 L petrol tank, 163 mm ground clearance. Concept schematic, not validated CAD.")
# schematic frame
fx, fy, fw, fh = MARGIN, y0, Inches(12.45), Inches(5.15)
rect(s, fx, fy, fw, fh, RGBColor(0xF8, 0xF9, 0xFA), line=GREY2)
legend_row(s, fx + Inches(0.15), fy + Inches(0.08), [("Added (new hydrogen parts)", GREEN), ("Retained", TAGBLUE), ("Modified", AMBER), ("Removed", GREYC)], size=11)
label(s, fx + fw - Inches(4.2), fy + Inches(0.08), Inches(4.1), Inches(0.32), "Side view, not to scale; dashed = control link", size=10.5, color=MIDGREY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
# car geometry (inches, front at left)
ox, ground = fx + Inches(1.45), fy + Inches(4.0)
def P(x, y): return (int(ox + Inches(x)), int(ground - Inches(y)))
body = [P(0.0, 0.6), P(0.2, 1.05), P(0.8, 1.38), P(2.5, 1.52), P(3.3, 2.2), P(6.1, 2.25), P(7.25, 1.6), P(8.85, 1.5), P(9.1, 1.1), P(9.15, 0.55), P(8.6, 0.35), P(0.5, 0.35)]
freeform(s, body, RGBColor(0xDC, 0xE7, 0xF2), line=NAVY2, line_width=1.5)
glass = [P(3.45, 1.45), P(3.95, 2.1), P(6.05, 2.13), P(7.0, 1.6), P(7.0, 1.45)]
freeform(s, glass, WHITE, line=NAVY2, line_width=1.0)
# door split line
line(s, *P(5.3, 1.45), *P(5.3, 2.12), color=NAVY2, width=0.75)
# ground line
line(s, *P(-0.6, 0.0), *P(9.8, 0.0), color=GREY2, width=1.0)
# zones text
label(s, *P(0.3, 1.95), Inches(1.6), Inches(0.28), "Engine bay", size=10, color=MIDGREY, italic=True)
label(s, *P(4.4, 1.3), Inches(1.2), Inches(0.28), "Cabin", size=10, color=MIDGREY, italic=True)
label(s, *P(7.55, 1.75), Inches(1.2), Inches(0.28), "Boot", size=10, color=MIDGREY, italic=True)
# components (car coordinates in inches; ground = 0)
def comp(x, y, w, h, fill, text=None, size=9, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3, line_col=None, dash=False, anchor=MSO_ANCHOR.MIDDLE):
    shp = rect(s, *P(x, y), Inches(w), Inches(h), fill, line=line_col, shape=shape, radius=radius)
    if dash and line_col is not None:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if text: shape_text(shp, text, size=size, color=WHITE, bold=True, margins=(0.02, 0, 0.02, 0.03), anchor=anchor)
    return shp
# lines first (they run under the floor and behind wheels)
line(s, *P(7.64, 0.97), *P(7.64, 0.28), color=GREEN, width=2.5)
line(s, *P(7.64, 0.28), *P(2.65, 0.28), color=GREEN, width=2.5)
line(s, *P(2.65, 0.28), *P(2.65, 0.93), color=GREEN, width=2.5)
line(s, *P(2.2, 0.9), *P(2.2, 0.2), color=GREYC, width=1.5); line(s, *P(2.2, 0.2), *P(4.2, 0.2), color=GREYC, width=1.5); line(s, *P(5.1, 0.2), *P(9.0, 0.2), color=GREYC, width=1.5)
for wx in (1.6, 7.6):
    dot(s, *P(wx, 0.45), Inches(0.95), RGBColor(0x2E, 0x2E, 0x2E)); dot(s, *P(wx, 0.45), Inches(0.45), RGBColor(0xC8, 0xCC, 0xD0))
comp(0.6, 1.5, 1.65, 0.62, TAGBLUE, "Engine 1.2 L, gearbox, catalyst (retained)", size=8, radius=0.1, anchor=MSO_ANCHOR.BOTTOM)
comp(1.2, 1.5, 0.85, 0.2, GREEN, "Injectors, rail", size=7.5, radius=0.5)
comp(0.7, 1.48, 0.3, 0.12, AMBER, None, radius=0.5)      # plugs / PCV marker
comp(2.4, 1.25, 0.5, 0.32, GREEN, "Reg.", size=8, radius=0.3)
comp(3.35, 1.95, 0.9, 0.3, GREEN, "H2 ECU", size=8.5, radius=0.4)
comp(5.75, 0.75, 1.5, 0.3, TAGBLUE, "Petrol tank 37 L (retained)", size=8, radius=0.4)
comp(7.7, 1.45, 1.35, 0.5, GREEN, "60 L H2 cylinder 350 bar", size=8.5, radius=0.5)
comp(7.57, 1.43, 0.14, 0.46, GREEN, None, radius=0.3)     # OTV
comp(8.1, 0.9, 1.0, 0.22, RGBColor(0xC4, 0xC7, 0xCA), "CNG 55 L removed", size=7, radius=0.5, line_col=GREYC, dash=True)
comp(8.72, 1.57, 0.2, 0.2, GREEN, None, shape=MSO_SHAPE.OVAL)   # receptacle
comp(8.2, 1.42, 0.15, 0.15, GREEN, None, shape=MSO_SHAPE.OVAL)  # boot sensor
comp(2.75, 1.55, 0.15, 0.15, GREEN, None, shape=MSO_SHAPE.OVAL)  # engine bay sensor
comp(4.2, 0.31, 0.9, 0.2, GREEN, "NOx trap", size=7.5, radius=0.5)
line(s, *P(7.32, 0.4), *P(7.32, 1.5), color=AMBER, width=3.0)          # bulkhead (modified)
line(s, *P(2.4, 1.1), *P(2.05, 1.4), color=GREEN, width=2.0)           # regulator -> injectors
line(s, *P(8.82, 1.37), *P(8.82, 1.45), color=GREEN, width=2.0)         # receptacle -> cylinder
line(s, *P(9.0, 0.95), *P(9.0, 0.08), color=GREEN, width=1.5, dash=True)   # TPRD vent
line(s, *P(3.35, 1.85), *P(2.05, 1.5), color=NAVY2, width=1.0, dash=True)    # ECU -> injectors
line(s, *P(4.25, 1.85), *P(7.57, 1.45), color=NAVY2, width=1.0, dash=True)   # ECU -> OTV
line(s, *P(8.2, 1.42), *P(4.25, 1.95), color=NAVY2, width=1.0, dash=True)    # boot sensor -> ECU
line(s, *P(2.83, 1.55), *P(3.35, 1.9), color=NAVY2, width=1.0, dash=True)   # bay sensor -> ECU
# leader labels on a grid: four above the car, five below
def leader(px_, py_, tx, ty, tw_, text, size=10, above=True):
    """label box at (tx, ty) in car coordinates (inches), width tw_; leader from the box edge to the component point."""
    bx_, by_ = P(tx, ty)
    label(s, bx_, by_, Inches(tw_), Inches(0.75), text, size=size, color=DARKGREY)
    ax_ = bx_ + Inches(tw_ / 2)
    ay_ = by_ + Inches(0.75) if above else by_
    line(s, ax_, ay_, *P(px_, py_), color=DARKGREY, width=0.75)
    dot(s, *P(px_, py_), Inches(0.07), DARKGREY)
topy = 3.5   # top label row (car coordinates, inches above ground)
leader(0.85, 1.48, -1.2, topy, 2.7, "Cold plugs, crankcase ventilation and lean calibration with retarded spark (modified)")
leader(3.8, 1.95, 1.75, topy, 2.7, "Hydrogen ECU: dual map, lambda 1.6 to 2.2, automatic petrol fallback, diagnostics (added)")
leader(8.27, 1.5, 4.7, topy, 2.7, "Twin hydrogen sensors in boot and engine bay; vent fan, cabin warning (added)")
leader(8.5, 1.45, 7.65, topy, 3.0, "60 L composite cylinder (OD 400 mm, 767 mm, 36 kg, 1.45 kg H2) with on-tank valve and TPRD; replaces the 55 L steel CNG cylinder; frame to 20 g / 8 g (added)", size=9.5)
boty = -0.18   # bottom label row (below ground)
leader(1.2, 0.9, -1.2, boty, 2.05, "Engine, gearbox and catalyst retained; injectors and rail on the intake (added)", above=False)
leader(2.65, 0.93, 1.05, boty, 2.05, "Two-stage regulator, 350 bar to injection pressure (added)", above=False)
leader(4.65, 0.3, 3.3, boty, 2.05, "316L high-pressure line under the floor; lean NOx trap provision on the exhaust (added)", above=False)
leader(6.5, 0.45, 5.55, boty, 2.05, "Petrol tank and injection retained, 555 km backup; boot bulkhead sealed and vented (modified)", above=False)
leader(9.0, 0.3, 7.8, boty, 2.05, "H35 receptacle with check valve (added); TPRD vent routed outside (added)", above=False)
# ---- right column: trade-off schematic chart
rx = fx + fw + Inches(0.25); rw = SLIDE_W - MARGIN - rx
cy = section_title(s, rx, y0, rw, "The central trade-off: lean burn buys low NOx at the cost of power")
lam = ["1.0", "1.3", "1.6", "2.0", "2.4", "2.8"]
power = (0.80, 0.72, 0.62, 0.50, 0.42, 0.36)
nox = (1.0, 2.7, 1.2, 0.15, 0.03, 0.0)
ch = line_chart(s, rx, cy, rw, Inches(2.4), lam, [("Power vs petrol (fraction)", power), ("Engine-out NOx vs stoichiometric (ratio)", nox)], [GREEN, REDC], number_format='0.0', font_size=10.5, legend=True, min_val=0, max_val=3.0)
ch.category_axis.has_title = True; ch.category_axis.axis_title.text_frame.text = "Air-fuel ratio, lambda"; ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10.5)
label(s, rx, cy + Inches(2.42), rw, Inches(1.0), "Schematic curves through literature anchors, not measured data: port injection loses about 20 percent power at lambda 1 and a further 35 to 40 percent at lambda 2 [W1-S9a, S9b]; NOx peaks near lambda 1.3 at about 2.7 times the stoichiometric value and falls to near zero above lambda 2 to 2.7 [W1-S9h, S9b]. Accelerations leave the lean window, hence the NOx trap provision [W4-S70].", size=10, color=MIDGREY)
cy2 = cy + Inches(3.5)
cy2 = section_title(s, rx, cy2, rw, "Packaging consequences (concept, mock-up in Phase 0)")
pk = [["Item", "CNG today", "Hydrogen", "Effect"], ["Cylinder OD x length", "316 x ~900 mm", "400 x 767 mm", "Boot loses ~50 L more"], ["Cylinder mass", "~55 kg steel", "36 kg composite", f"Kerb mass {E['netmass']:+.0f} kg"], ["Fuel on board", "~9 kg CNG", f"{E['usable']:.2f} kg usable H2", "35 percent of energy"], ["Seats / luggage", "5 / ~200 L", "5 / ~150 L", "Two bags become one"]]
table2(s, rx, cy2, rw, Inches(1.5), pk, col_widths=[1.3, 1.1, 1.15, 1.35], font_size=10.5, header_size=10.5, row_heights=[Inches(0.3)] + [Inches(0.3)] * 4)
cy3 = section_title(s, rx, cy2 + Inches(1.65), rw, "Consumption cross-check (kg H2 per 100 km)")
xc = [["Vehicle", "kg/100 km", "Basis"], ["This retrofit, fleet use (model)", f"{IN['h2_cons']['active']:.2f}", f"0.87 x CNG efficiency at 25 km/kg"], ["Ford Focus H2RV (2.3 L, hybrid)", "1.4", "claimed, 2003 [W1-S12c]"], ["Mazda RX-8 Hydrogen RE, rotary", "2.4", "claimed, 2006 [W1-S12b]"], ["Toyota hydrogen HiAce (3.4 L van)", "3.1", "6.2 kg for 200 km [W1-S12e]"]]
table2(s, rx, cy3, rw, Inches(1.45), xc, col_widths=[2.3, 0.8, 1.9], font_size=10, header_size=10, row_heights=[Inches(0.28)] + [Inches(0.29)] * 4)
# ---- spec strip
sy = fy + fh + Inches(0.15)
specs = [(f"{E['usable']:.2f} kg", "usable H2 per fill (1.44 kg nominal, 85 percent fill, 20 bar residual)"), (f"{E['range']:.0f} km", f"per fill at {IN['h2_cons']['active']:.2f} kg/100 km; {E['range700']:.0f} km at 700 bar"),
         (f"~{E['usable']*2/(IN['h2_cons']['active']/100):.0f} km", f"per day on hydrogen with two depot fills ({E['maxshare']*100:.0f} percent of a 167 km day)"), (f"{E['h2ps']:.0f} PS", f"peak in hydrogen mode, {100*(1-IN['h2_power_factor']['active']):.0f} percent below petrol; petrol switch at high load"),
         (f"{E['netmass']:+.0f} kg", "net kerb mass; payload unchanged"), ("~150 L", "boot left (about 200 L with CNG); seats unchanged"), ("3 days", "conversion downtime per car (14 labour hours at volume)")]
kw_ = (fw - Inches(0.9)) / 7
for i, (v, c) in enumerate(specs):
    kpi(s, fx + i * (kw_ + Inches(0.15)), sy, kw_, Inches(1.2), v, c, value_size=17, cap_size=9.5)
# bottom-left: what cannot be reused
by = sy + Inches(1.35)
cy = section_title(s, fx, by, fw, "What cannot be reused, and the injection choice")
label(s, fx, cy, fw / 2 - Inches(0.1), Inches(1.1), "Removed, not reusable: the 55 L steel CNG cylinder (ISO 11439 steels are limited to 2 percent hydrogen), CNG regulator, rail, injectors, hoses and lines. Hydrogen parts need UN R134 or ISO 19881 class qualification [W1-S11a].", size=11, color=DARKGREY)
label(s, fx + fw / 2 + Inches(0.1), cy, fw / 2 - Inches(0.1), Inches(1.1), "Port injection timed after intake-valve closing prevents backfire and pre-ignition but costs power; direct injection (Suzuki-AVL Swift prototype) recovers it but is not a retrofit. Durability to prove: valve seats and guides, lubricant ash, sensor drift.", size=11, color=DARKGREY)
footer2(s, "Sources: Maruti Suzuki Tour S and Dzire specifications [W1-S3a to S3c]; Luxfer G-Stor Pro V060H35 [W1-S2]; hydrogen densities [W1-S13a]; PFI power loss and NOx behaviour [W1-S9a, S9b, S9h]; ISO 11439 limit [W1-S11a]; UN R110 loads [W1-S14d]; Energy_Range and BOM sheets.", 4)

# =============================================================== SLIDE 5
s = blank_slide(prs)
y0 = header2(s, "Safety and approval: each hazard has a known prevention, detection and test; the unresolved item is approval, because no route exists yet for a hydrogen retrofit of a car",
             "Hydrogen: flammable from 4 to 75 percent in air at 0.017 mJ, diffuses 3.8 times faster than natural gas, burns almost invisibly [W2-S43, S44]. Different risks, not automatically higher, if designed for.")
tw = Inches(12.4)
rows = [["Hazard", "Prevention (design)", "Detection and response", "Evidence required", "Evidence type"],
 ["Leak into boot, cabin, garage", "Sealed bulkhead, vented boot, no ignition sources in the bay, valve closed with ignition off", "Two sensors (1 and 2 percent), tank valve closes within 1 s, switch to petrol, cabin warning", "Below 4 percent in cabin, below 8 percent at vents (GTR 13 / AIS-195); parked enclosed-space test", "Type test"],
 ["Embrittlement, permeation, ageing", "316L lines, hydrogen elastomers, composite liner; no CNG steel parts reused", "Leak check at service; 3-yearly requalification; 15-year life", "ISO 19881 / UN R134 supplier qualification; cycling and permeation", "Standard"],
 ["Crash damage", "Cylinder inside the crush-free zone; frame to 20 g / 8 g", "Inertia switch closes the tank valve; post-crash leak limit", "Sled test; full-vehicle rear crash under AIS-195", "Type test"],
 ["Overpressure and fire", "TPRD vents away from occupants and wheel wells; burst ratio 2.25 x NWP", "TPRD opens on heat: controlled release, not rupture", "Bonfire, burst, drop, cycling (GTR 13); vent-routing review", "Type test"],
 ["Detection or control failure", "Redundant sensors, ECU watchdog, manual shut-off, limp-home on petrol", "Fault codes, forced petrol mode, service lock-out", "Fault injection and hardware-in-the-loop; 500 h engine durability", "Programme test"],
 ["Workshop and refuelling", "NFPA 2 ventilation (1 cfm per sq ft), detection, permit-to-work, PESO-trained staff", "Dispenser and bay interlocks; nitrogen purge for venting", "PESO Form-H licence; SAE J2601 fill protocol; incident drills", "Licence"],
 ["Combustion emissions (NOx)", "Lean burn lambda 1.6 to 2.2, retarded spark, NOx trap provision, load limiting", "Lambda sensor and OBD", "BS-VI 60 mg/km on the certification cycle; quarterly pilot tests", "Target"],
]
table2(s, MARGIN, y0, tw, Inches(6.5), rows, col_widths=[1.35, 2.6, 2.45, 2.7, 1.0], font_size=10.5, header_size=11.5)
rx = MARGIN + tw + Inches(0.25); rw = SLIDE_W - MARGIN - rx
cy = section_title(s, rx, y0, rw, "Approval sequence: what exists and what is missing")
steps = [("New-vehicle hydrogen ICE standard", "AIS-195:2023 (Amd 1, May 2025) via CMVR Rule 125M, GSR 746(E) [W2-S13]", "Established"),
         ("Retrofit gateway in law", "MV Act s.52; Rule 126 proviso; Rule 115 kit machinery written for CNG and LPG [W2-S3, S9]", "Established"),
         ("Hydrogen retrofit kit standard and emission procedure", "None; AIS-157 excludes retrofitted FCEVs; AIS-195 retrofit position unverifiable [W2-S11]", "Gap"),
         ("Car-size cylinder approval (PESO)", "Gas Cylinder Rules 2025 cover on-board hydrogen; only 150 to 250 L truck cylinders approved [W2-S24, S30]", "Unresolved"),
         ("Dispensing licence", "PESO Form-H online since May 2025; SMPV(U) 2025 for bulk storage [W2-S27]", "Established"),
         ("Registration, insurance, warranty", "Form 22D endorsement needs an approved kit; IMT-25 covers CNG/LPG only; OEM warranty void [W2-S53, S54]", "Unresolved"),
         ("Trial route for the pilot", "Trade certificate plus test-agency evaluation (Mirai 2022, Tata trucks 2025); MNRE scheme names 4-wheelers", "Established")]
for i, (t, d_, st) in enumerate(steps):
    ry = cy + i * Inches(0.86)
    numbered_circle(s, rx, ry + Inches(0.04), Inches(0.36), i + 1, fill=STATUS[st], size=12)
    label(s, rx + Inches(0.48), ry, rw - Inches(1.9), Inches(0.34), t, size=11.5, color=DEEPNAVY, bold=True)
    label(s, rx + Inches(0.48), ry + Inches(0.31), rw - Inches(1.9), Inches(0.55), d_, size=10, color=DARKGREY)
    status_chip(s, rx + rw - Inches(1.3), ry + Inches(0.06), st, w=Inches(1.3), size=10)
    if i < 6: line(s, rx + Inches(0.18), ry + Inches(0.42), rx + Inches(0.18), ry + Inches(0.9), color=GREY2, width=1.0)
# bottom strip
by = y0 + Inches(6.7)
callout(s, MARGIN, by, tw, Inches(1.5), "NOx is the emission that matters: hydrogen combustion still forms thermal NOx (US DOE [W2-S52]). Lean burn above lambda 2 is near zero at steady state, but accelerations leave that window, so the kit carries a lean NOx trap provision and gate 2 requires 60 mg/km or better on the BS-VI cycle. Target, not a result. No 'approved' or 'fully safe' claim is made; safety evidence comes from type tests and standards, not from incident-free pilot operation.", fill=YELLOW, size=11.5)
callout(s, rx, by, rw, Inches(1.5), "Gate 1 (March 2027): a written regulator and test-agency view on the retrofit route, or the programme stops. Legend: green = in force; amber = exists but not for this case; red = missing instrument.", fill=MINT, size=11.5, bold=True, color=DARKGREEN)
footer2(s, "Sources: H2Tools and DOE properties [W2-S43, S44]; NFPA 2 [W2-S45]; GTR 13 limits [W2-S42]; GSR 746(E) [W2-S13]; Rules 126 and 115 [W2-S9, S3]; AIS-157 scope [W2-S11]; Gas Cylinder Rules 2025 [W2-S24]; Form-H [W2-S27]; IMT-25, Form 22D [W2-S53, S54]; incidents [W2-S48, S49].", 5)

# =============================================================== SLIDE 6
s = blank_slide(prs)
y0 = header2(s, "Economics: over five years the converted car costs the owner ₹16.8 lakh more than staying on CNG and ₹16.2 lakh more than switching to a BEV; fuel is the whole story",
             "Fair comparison for an owner deciding in 2028 with a 3-year-old Dzire Tour S CNG. Same horizon, mileage, discount rate and time costs; donor sale credited to the BEV option; residual values credited to all.")
T = CD["tco_cats"]
def cat(opt, keys): return sum(T[opt].get(k, 0) for k in keys) / 1e5
capA = T["A"]["cap"] + T["A"]["resid"]; capB = T["B"]["cap"] + T["B"]["resid"]; capC = T["C"]["cap"] + T["C"]["resid"]
fuelA = cat("A", ["CNG fuel", "Petrol fuel (share of km)"]); fuelB = cat("B", ["Hydrogen fuel", "Petrol fuel (backup share of km)"]); fuelC = cat("C", ["Electricity (depot and public mix)"])
mA, mB, mC = cat("A", ["Maintenance and tyres"]), cat("B", ["Maintenance and tyres"]), cat("C", ["Maintenance and tyres"])
iA = cat("A", ["Insurance", "Cylinder periodic test (annualised)"]); iB = cat("B", ["Insurance", "Hydrogen system inspection and cylinder requalification"]); iC = cat("C", ["Insurance"])
tA = cat("A", ["In-shift refuelling time cost"]); tB = cat("B", ["In-shift refuelling time and detour cost"]); tC = cat("C", ["In-shift charging time cost"])
cats_ = [f"Continue on CNG\n₹{S['npvA']/1e5:.1f} lakh, ₹{S['lcA']:.1f}/km", f"Hydrogen retrofit\n₹{S['npvB']/1e5:.1f} lakh, ₹{S['lcB']:.1f}/km", f"Sell, buy compact BEV\n₹{S['npvC']/1e5:.1f} lakh, ₹{S['lcC']:.1f}/km"]
series = [("Capital net of residual value", (round(capA/1e5, 2), round(capB/1e5, 2), round(capC/1e5, 2))), ("Fuel or energy", (round(fuelA, 2), round(fuelB, 2), round(fuelC, 2))),
          ("Maintenance and tyres", (round(mA, 2), round(mB, 2), round(mC, 2))), ("Insurance and inspection", (round(iA, 2), round(iB, 2), round(iC, 2))), ("In-shift time cost", (round(tA, 2), round(tB, 2), round(tC, 2)))]
lx, lw = MARGIN, Inches(7.9)
cy = section_title(s, lx, y0, lw, "Five-year cost of ownership by category, ₹ lakh NPV (model, base case)")
ch = stacked_chart(s, lx, cy, Inches(5.1), Inches(5.2), cats_, series, [NAVY2, GREEN, CHARTBLUE, MIDGREY, AMBER], number_format='0.0', font_size=10, label_size=10, value_axis_title="₹ lakh, present value at 12 percent", data_labels=False, legend=False)
for si in (0, 1):
    dl = ch.plots[0].series[si].data_labels; dl.show_value = True; dl.number_format = '0.0'; dl.number_format_is_linked = False; dl.font.size = Pt(10); dl.font.color.rgb = WHITE; dl.font.bold = True
    from pptx.enum.chart import XL_LABEL_POSITION as _XLP; dl.position = _XLP.CENTER
ctab = [["₹ lakh", "CNG", "H2", "BEV"], ["Capital net of residual", f"{capA/1e5:.1f}", f"{capB/1e5:.1f}", f"{capC/1e5:.1f}"], ["Fuel or energy", f"{fuelA:.1f}", f"{fuelB:.1f}", f"{fuelC:.1f}"], ["Maintenance, tyres", f"{mA:.1f}", f"{mB:.1f}", f"{mC:.1f}"], ["Insurance, inspection", f"{iA:.1f}", f"{iB:.1f}", f"{iC:.1f}"], ["In-shift time", f"{tA:.1f}", f"{tB:.1f}", f"{tC:.1f}"], ["Total NPV", f"{S['npvA']/1e5:.1f}", f"{S['npvB']/1e5:.1f}", f"{S['npvC']/1e5:.1f}"], ["₹ per km", f"{S['lcA']:.1f}", f"{S['lcB']:.1f}", f"{S['lcC']:.1f}"]]
tbc = table2(s, lx + Inches(5.2), cy + Inches(0.3), Inches(2.7), Inches(3.4), ctab, col_widths=[1.5, 0.6, 0.6, 0.6], font_size=10, header_size=10, bold_rows=(6, 7), row_heights=[Inches(0.3)] * 8)
for ri, col in enumerate([NAVY2, GREEN, CHARTBLUE, MIDGREY, AMBER], start=1):
    c_ = tbc.cell(ri, 0); c_.fill.solid(); c_.fill.fore_color.rgb = col
    for r_ in c_.text_frame.paragraphs[0].runs: r_.font.color.rgb = WHITE
label(s, lx + Inches(5.2), cy + Inches(3.8), Inches(2.7), Inches(1.3), "Colours match the chart segments. Capital is net of residual value (negative for the kept CNG car). Petrol backup km are inside 'fuel or energy' for the retrofit (₹5.0 lakh).", size=9.5, color=MIDGREY)
label(s, lx, cy + Inches(5.25), lw, Inches(0.5), f"Fuel or energy: hydrogen km ₹{S['fkm_h2']:.1f}/km at ₹{IN['h2_price_y1']['active']:.0f}/kg falling 10 percent a year, petrol backup km ₹{S['fkm_petrol']:.1f}, CNG ₹{S['fkm_cng']:.1f}, BEV ₹{S['fkm_bev']:.1f}. Negative capital for the CNG car is its residual value.", size=10.5, color=MIDGREY)
# middle: conversion cost breakdown
mx, mw = lx + lw + Inches(0.3), Inches(5.3)
cy = section_title(s, mx, y0, mw, "Conversion cost by category, ₹ lakh installed (BOM sheet)")
bom = CD["bom"]
def bsum(idx, col): return sum(bom[i][col] for i in idx) / 1e5
groups = [("Storage and mounting", [0, 1, 2]), ("Fuel delivery", [3, 4, 5, 6, 7]), ("Controls and safety", [8, 9]), ("Engine and NOx provision", [10, 11]), ("Installation and testing", [12, 13]), ("Certification allocation", [14])]
gp = [(g, bsum(ix, 1), bsum(ix, 2)) for g, ix in groups]
rest_p = B["pilot"] / 1e5 - sum(v for _, v, _ in gp); rest_v = B["volume"] / 1e5 - sum(v for _, _, v in gp)
gp.append(("Warranty, margin, GST", rest_p, rest_v))
cser = [(g, (round(p_, 2), round(v_, 2))) for g, p_, v_ in gp]
stacked_chart(s, mx, cy, mw, Inches(3.0), [f"Pilot batch\n₹{L(B['pilot'])} lakh", f"10,000 per year\n₹{L(B['volume'])} lakh"], cser, [GREEN, TAGBLUE, NAVY2, AMBER, CHARTBLUE, MIDGREY, GREY2], number_format='0.00', font_size=9.5, legend=True, data_labels=False, value_axis_title="₹ lakh")
crow = [["Category", "Pilot", "Volume"]] + [[g, f"{p_:.2f}", f"{v_:.2f}"] for g, p_, v_ in gp] + [["Installed total incl. GST", f"{B['pilot']/1e5:.2f}", f"{B['volume']/1e5:.2f}"]]
table2(s, mx, cy + Inches(3.1), mw, Inches(2.3), crow, col_widths=[2.3, 0.8, 0.8], font_size=10, header_size=10.5, bold_rows=(8,), row_heights=[Inches(0.26)] * 9)
# right: heat grid + break-even
rx, rw = mx + mw + Inches(0.3), SLIDE_W - MARGIN - (mx + mw + Inches(0.3))
cy = section_title(s, rx, y0, rw, "Extra cost vs staying on CNG, ₹ lakh over 5 years")
g1 = O["SENS"]["grid1"]; kmcols = [1, 3, 5]
sel = [r for r in g1[1:] if r[0] in (150, 250, 400, 600, 800)]
rows = [["₹/kg \\ km/yr"] + [n0(g1[0][c]) for c in kmcols]] + [[f"₹{r[0]:.0f}"] + [f"+{r[c]:.1f}" for c in kmcols] for r in sel]
tb = table2(s, rx, cy, rw, Inches(2.4), rows, col_widths=[1.2, 1, 1, 1], font_size=11, header_size=10.5, row_heights=[Inches(0.36)] * 6)
def heat(vv):
    t = max(0, min(22, vv)) / 22
    return RGBColor(int(0xC3 + (0xE0 - 0xC3) * t), int(0xED - (0xED - 0x7A) * t), int(0xD9 - (0xD9 - 0x66) * t))
for i, r in enumerate(sel, start=1):
    for j, c in enumerate(kmcols, start=1):
        cell = tb.cell(i, j); cell.fill.solid(); cell.fill.fore_color.rgb = heat(r[c])
label(s, rx, cy + Inches(2.45), rw, Inches(0.5), "Every cell is positive: no cell turns negative in the workbook grid down to ₹100/kg. More mileage widens the gap while hydrogen is above ₹188/kg.", size=10, color=MIDGREY)
cy2 = section_title(s, rx, cy + Inches(3.05), rw, "Break-even indicators (model)")
be = [("CNG fuel parity", f"₹{S['parity_cng']:.0f}/kg delivered"), ("Petrol fuel parity", f"₹{S['parity_petrol']:.0f}/kg"), ("Full-TCO break-even vs CNG", "none in 5 years at any price"), ("Full-TCO break-even vs BEV", "none at base; needs about ₹150/kg, ₹2 lakh kit, 85 percent hydrogen km"), ("Abatement cost vs CNG", f"₹{EM['abate']/1000:.0f},000/t CO2e at base; about ₹45,000/t at ₹250/kg (credits ₹1,350 to 2,800/t)")]
kv_table(s, rx, cy2, rw, be, kw=1.9, size=10.5, row_h=Inches(0.46))
# assumptions strip
ay = y0 + Inches(6.2)
cy = section_title(s, MARGIN, ay, SLIDE_W - 2 * MARGIN, "Assumptions used in every exhibit on this slide (Inputs sheet; downside and upside cases in the annexure)")
chips = [f"Decision 2028, 5 years", f"{IN['km_day']['active']*IN['days']['active']:,.0f} km/yr", f"Discount 12 percent", f"H2 ₹{IN['h2_price_y1']['active']:.0f}/kg, minus 10 percent/yr", f"CNG ₹{IN['cng_price']['active']:.0f}/kg at {IN['cng_kmkg']['active']:.0f} km/kg", f"Petrol ₹{IN['petrol_price']['active']:.0f}/L", f"H2 {IN['h2_cons']['active']:.2f} kg/100 km, {IN['h2_share']['active']*100:.0f} percent of km", f"Power ₹{IN['depot_tariff']['active']:.0f} depot, ₹{IN['pub_tariff']['active']:.0f} public per kWh", f"BEV ₹{L(IN['bev_price']['active'])} lakh, resale ₹{L(IN['bev_resid']['active'])} lakh", f"Donor sale ₹{L(IN['donor_value_now']['active'])} lakh", "Time ₹150/h in shift"]
cx = MARGIN; cyy = cy
for c in chips:
    w_ = Inches(0.105 * len(c) + 0.35)
    if cx + w_ > SLIDE_W - MARGIN:
        cx = MARGIN; cyy += Inches(0.44)
    chip(s, cx, cyy, w_, Inches(0.36), c, RGBColor(0xEE, 0xF1, 0xF0), color=DARKGREY, size=10.5, bold=False); cx += w_ + Inches(0.12)
callout(s, MARGIN, cyy + Inches(0.55), SLIDE_W - 2 * MARGIN, Inches(0.95), f"Reading: the retrofit is a decarbonisation cost, not a saving. Hydrogen km cost ₹{S['fkm_h2']:.1f} against ₹{S['fkm_cng']:.1f} on CNG, and the {100-IN['h2_share']['active']*100:.0f} percent of km on petrol backup add ₹{cat('B', ['Petrol fuel (backup share of km)']):.1f} lakh. It needs a paying counterparty or a cheaper molecule than any 2030 forecast found (CEEW ₹360/kg delivered; refinery-gate record ₹279/kg before compression).", fill=MINT, size=11.5, bold=True, color=DARKGREEN)
footer2(s, "Sources: fuel prices Sep 2026 [W4-S16, S19]; DERC tariff, public charging [W4-S20, S21]; BEV prices, tests, warranties [W3-S27, S30; W4-S26]; NITI Aayog mileage [W4-S30]; CEEW 2025 [W4-S22]; Numaligarh tender [W3-S39]; carbon credits [W4-S73]; TCO, BOM and Sensitivities sheets.", 6)

# =============================================================== SLIDE 7
s = blank_slide(prs)
y0 = header2(s, f"Hydrogen supply: a 30-car pilot needs about {bd['demand'][0]:.0f} kg a day; delivered hydrogen costs about ₹{bd['delivered'][0]:,.0f}/kg at a pilot point and ₹{bd['delivered'][1]:,.0f}/kg at a 2030 hub, and utilisation alone cannot reach parity",
             "Pilot setting: Faridabad, Delhi-NCR (IOCL R&D dispensing since the 2022 Mirai pilot; MNRE corridor sanctioned March 2025, 18 to 24 months to commission). No station is public or confirmed available: access is a proposed partnership.")
# flow
cy = section_title(s, MARGIN, y0, SLIDE_W - 2 * MARGIN, "Supply-to-fleet flow with pilot quantities (model, H2_Supply sheet) and proposed owners")
nodes = [("Green hydrogen producer", f"SIGHT-awarded plant; ₹{bd['prod'][0]:.0f}/kg at gate (2026-28), ₹{bd['prod'][1]:.0f}/kg (2030); certified 2 kg CO2e/kg", "Producer"),
         ("Tube-trailer delivery", f"300 kg payload; {bd['deliv'][0]:.1f} loads per week; ₹{bd['trans'][0]:.0f}/kg; about 150 km", "Logistics contractor"),
         ("Cascade storage", f"{bd['buffer'][0]:.0f} kg buffer (1.5 days) at 450 bar", "OMC station partner"),
         ("Compression", f"{bd['comp'][0]:.1f} kg/h on a 150 kg/day design; {IN['comp_kwh']['active']:.0f} kWh/kg", "OMC station partner"),
         ("Dispensing (H35)", f"{bd['disp'][0]:.0f} positions; {bd['fills'][0]:.0f} fills/day, 10 per hour in the peak window; PESO Form-H", "OMC station partner"),
         ("Fleet consumption", f"30 cars x 2 fills; {bd['demand'][0]:.0f} kg/day; {IN['h2_share']['active']*100:.0f} percent of km on hydrogen", "Fleet partner")]
nw = (SLIDE_W - 2 * MARGIN - 5 * Inches(0.35)) / 6; nh = Inches(1.45)
for i, (t, d_, own) in enumerate(nodes):
    nx = MARGIN + i * (nw + Inches(0.35))
    rect(s, nx, cy, nw, nh, WHITE, line=GREY2); rect(s, nx, cy, nw, Inches(0.38), NAVY2 if i < 5 else GREEN)
    label(s, nx + Inches(0.08), cy, nw - Inches(0.1), Inches(0.38), f"{i+1}. {t}", size=11.5, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    label(s, nx + Inches(0.08), cy + Inches(0.42), nw - Inches(0.16), Inches(0.8), d_, size=10.5, color=DARKGREY)
    chip(s, nx + Inches(0.08), cy + nh - Inches(0.36), nw - Inches(0.16), Inches(0.28), own, RGBColor(0xEE, 0xF1, 0xF0), color=DEEPNAVY, size=9.5)
    if i < 5: arrow(s, nx + nw, cy + Inches(0.8), nx + nw + Inches(0.35), cy + Inches(0.8), color=GREEN, width=2.5)
# utilisation chart
uy = cy + nh + Inches(0.2)
lx, lw = MARGIN, Inches(9.2)
cy2 = section_title(s, lx, uy, lw, "Delivered price falls with utilisation but never reaches CNG parity (model)")
ut = H["util_table"]; ucat = [f"{int(u*100)} percent" for u, _, _ in ut]
line_chart(s, lx, cy2, lw, Inches(2.65), ucat, [("Pilot point 2028: 150 kg/day, delivered H2", tuple(round(p_) for _, p_, _ in ut)), ("Hub 2030: 500 kg/day, delivered H2", tuple(round(h_) for _, _, h_ in ut)), (f"Petrol fuel parity ₹{S['parity_petrol']:.0f}/kg", tuple([round(S['parity_petrol'])] * 6)), (f"CNG fuel parity ₹{S['parity_cng']:.0f}/kg", tuple([round(S['parity_cng'])] * 6))], [NAVY2, GREEN, AMBER, REDC], number_format='#,##0', font_size=10.5, legend=True, value_axis_title="₹ per kg delivered", min_val=0, data_labels=False)
label(s, lx, cy2 + Inches(2.67), lw, Inches(0.6), f"At 95 percent utilisation the hub price is still ₹{ut[-1][2]:,.0f}/kg: production, transport and compression alone (₹{bd['prod'][1]+bd['trans'][1]+bd['elec'][1]:.0f}/kg) exceed CNG parity. Station utilisation is the axis; capex ₹{IN['station_capex_pilot']['active']/1e7:.0f} crore (pilot) and ₹{IN['station_capex_2030']['active']/1e7:.0f} crore (hub) recovered over 10 years at 12 percent.", size=10.5, color=MIDGREY)
# responsibility table
rx, rw = lx + lw + Inches(0.3), SLIDE_W - MARGIN - (lx + lw + Inches(0.3))
cy3 = section_title(s, rx, uy, rw, "Ownership and risk (proposed, not committed)")
rows = [["Function", "Proposed holder", "Risk carried"], ["Vehicles", "Fleet partner (PSU, corporate or airport contract)", "Residual value, backed by programme buy-back"], ["Kits", "SRDI-led consortium; kit integrator warrants", "Development and warranty"], ["Dispensing point", "OMC partner (proposed IOCL, Faridabad), Form-H licence", "Utilisation, covered by take-or-pay from the fleet"], ["Fuel", "SIGHT-awarded producer by tube trailer", "Supply reliability; buffer covers 1.5 days"], ["Service", "Two hydrogen-ready workshops; PESO test station", "Safety compliance and training"], ["Evaluation", "IIT Delhi or CEEW (proposed)", "Credibility of results"]]
table2(s, rx, cy3, rw, Inches(2.7), rows, col_widths=[1.1, 2.3, 2.1], font_size=10.5, header_size=11, row_heights=[Inches(0.3)] + [Inches(0.4)] * 6)
pby = cy2 + Inches(3.3)
cyp = section_title(s, lx, pby, lw, "Delivered price build-up, ₹ per kg (model): pilot point 2028 versus 2030 hub")
lm = [bd['delivered'][i] - (bd['prod'][i] + bd['trans'][i] + bd['elec'][i] + bd['caprec'][i] + bd['om'][i]) for i in range(2)]
prow = [["Component", "Pilot point, 150 kg/day at 50 percent", "Hub 2030, 500 kg/day at 70 percent"], ["Production at plant gate (green)", f"{bd['prod'][0]:.0f}", f"{bd['prod'][1]:.0f}"], ["Tube-trailer transport", f"{bd['trans'][0]:.0f}", f"{bd['trans'][1]:.0f}"], ["Compression electricity", f"{bd['elec'][0]:.0f}", f"{bd['elec'][1]:.0f}"], ["Station capital recovery (10 years, 12 percent)", f"{bd['caprec'][0]:.0f}", f"{bd['caprec'][1]:.0f}"], ["Station fixed O&M", f"{bd['om'][0]:.0f}", f"{bd['om'][1]:.0f}"], ["Losses, margin and taxes", f"{lm[0]:.0f}", f"{lm[1]:.0f}"], ["Delivered price at the dispenser", f"{bd['delivered'][0]:,.0f}", f"{bd['delivered'][1]:,.0f}"]]
table2(s, lx, cyp, lw, Inches(2.05), prow, col_widths=[2.6, 1.7, 1.7], font_size=10.5, header_size=10.5, bold_rows=(7,), row_heights=[Inches(0.3)] + [Inches(0.25)] * 7, align_center_from=1)
callout(s, rx, cy3 + Inches(2.85), rw, Inches(1.0), "Main bottleneck: not station throughput (2 dispensers cover 60 fills a day) but the hydrogen price stack. Delivered gas keeps pilot capital to a dispensing point; on-site electrolysis is justified only at hub scale. Petrol backup covers supply interruptions.", fill=YELLOW, size=11)
cys = section_title(s, rx, cy3 + Inches(4.0), rw, "Station sizing from the fleet (H2_Supply sheet)")
sz = [f"30 cars x 167 km x {IN['h2_cons']['active']/100:.4f} kg/km x {IN['h2_share']['active']:.2f} = {bd['demand'][0]:.0f} kg/day", f"{bd['fills'][0]:.0f} fills/day, 50 percent in a 3 h window = 10/h, 2 dispensers", f"Compressor {bd['comp'][0]:.1f} kg/h over 16 h; buffer {bd['buffer'][0]:.0f} kg (1.5 days)", f"Cars alone use {bd['demand'][0]/bd['cap'][0]*100:.0f} percent of capacity; 50 percent assumed with bus and truck users", f"One 2030 hub serves {bd['carsper'][1]:.0f} cars at 70 percent utilisation"]
textbox(s, rx, cys, rw, Inches(1.6), [(t, {'bullet': True}) for t in sz], size=10.5, color=DARKGREY, line_spacing=1.03, space_after=2)
footer2(s, "Sources: MNRE corridors, PIB 3 Mar 2025 [W3-S36]; IOCL Faridabad [W3-S46, W4-S81]; CEEW station model [W4-S22]; ICCT utilisation [W4-S46]; Indian station capex indication [W4-S56]; green hydrogen cost [W4-S33, S38, S57]; Gas Cylinder Rules, Form-H [W2-S24, S27]; H2_Supply sheet.", 7)

# =============================================================== SLIDE 8
s = blank_slide(prs)
y0 = header2(s, f"Roadmap to 2030: four gated phases across five workstreams, ₹{tot_cr:.1f} crore including contingency, with a proceed, redesign or stop decision at each gate",
             "October 2026 to September 2030. Suzuki R&D Center India proposed as engineering lead; public support sought under the MNRE transport pilot scheme, which names buses, trucks and 4-wheelers (PIB 14 Feb 2024).")
gx = MARGIN + Inches(2.55); gw = SLIDE_W - MARGIN - gx; ly = y0
quarters = 17  # Q4 2026 .. Q4 2030
qw = gw / quarters
# quarter header
years = [("2026", 1), ("2027", 4), ("2028", 4), ("2029", 4), ("2030", 4)]
qx = gx
for yname, nq in years:
    rect(s, qx, ly, qw * nq, Inches(0.3), NAVY2); label(s, qx, ly, qw * nq, Inches(0.3), yname, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    qx += qw * nq
for q in range(quarters):
    line(s, gx + qw * q, ly + Inches(0.3), gx + qw * q, ly + Inches(5.55), color=GREY2, width=0.5)
# phase band (quarters: P0 = Q4'26-Q1'27 (2q), P1 = Q2'27-Q1'28 (4q), P2 = Q2'28-Q1'29 (4q), P3 = Q2'29-Q3'30 (6q), P4 = Q4'30 (1q))
phases = [("Phase 0: feasibility", 0, 2, NAVY2, ph_cr[0]), ("Phase 1: engineering and controlled validation", 2, 4, TAGBLUE, ph_cr[1]), ("Phase 2: certification and pilot build", 6, 4, GREEN, ph_cr[2]), ("Phase 3: pilot operation", 10, 6, DEEPGREEN, ph_cr[3]), ("P4", 16, 1, DARKGREEN, None)]
gy_ = ly + Inches(0.36)
label(s, MARGIN, gy_, Inches(2.5), Inches(0.36), "Decision gates", size=11, color=REDC, bold=True, anchor=MSO_ANCHOR.MIDDLE)
py = ly + Inches(0.78)
label(s, MARGIN, py, Inches(2.5), Inches(0.36), "Phase and budget (₹ crore)", size=11, color=DEEPNAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
for t, q0, nq, col, bud in phases:
    rect(s, gx + qw * q0, py, qw * nq - Inches(0.03), Inches(0.36), col)
    label(s, gx + qw * q0, py, qw * nq, Inches(0.36), f"{t}  ₹{bud:.1f} cr" if bud is not None else "Cond.", size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# gates
gates = [(2, "G1 Mar 2027"), (6, "G2 Mar 2028"), (10, "G3 Mar 2029"), (16, "G4 Sep 2030")]
for q, t in gates:
    diamond(s, gx + qw * q - Inches(0.17), gy_ + Inches(0.01), Inches(0.34), REDC)
    label(s, gx + qw * q + Inches(0.2), gy_, Inches(1.6), Inches(0.36), t, size=10, color=REDC, bold=True, anchor=MSO_ANCHOR.MIDDLE)
# lanes
lanes = [("Engineering", "SRDI lead", [(0, 2, "Packaging mock-up, CAD, hazard analysis"), (2, 4, "500 h test-cell programme; 3 prototype conversions"), (6, 4, "30 conversions; installer training"), (10, 6, "Field support, calibration updates")]),
         ("Testing and approval", "ARAI or ICAT; PESO", [(0, 2, "Route consultation: AISC, MoRTH, ARAI, ICAT, PESO"), (2, 4, "Storage qualification; leak, ventilation, crash-structure tests"), (6, 4, "Kit type approval; 80,000 km durability on 2 cars"), (10, 6, "Quarterly emission and safety tests")]),
         ("Fuel supply", "OMC partner", [(0, 2, "Station partner letter; supply options"), (2, 4, "Supply contract; Form-H application"), (6, 4, "Dispensing point commissioned"), (10, 6, "Deliveries 1.4 loads/week; utilisation log")]),
         ("Fleet operations", "Fleet partner", [(0, 2, "30 interviews; telematics on 20 cars"), (2, 4, "Driver and workshop training plan"), (6, 4, "2 workshops upgraded; 30 cars inducted"), (10, 6, "18 months with 10 CNG and 10 BEV comparison cars")]),
         ("Commercial validation", "SRDI with evaluator", [(0, 2, "Demand evidence from 3 fleets"), (2, 4, "Cost model with supplier quotations"), (6, 4, "MNRE sanction; insurer endorsement"), (10, 6, "Independent evaluation; scale decision")])]
lh = Inches(0.86); ly2 = py + Inches(0.45)
for i, (name, owner, bars) in enumerate(lanes):
    yy = ly2 + i * lh
    if i % 2 == 0: rect(s, MARGIN, yy, SLIDE_W - 2 * MARGIN, lh, RGBColor(0xF6, 0xF7, 0xF8))
    label(s, MARGIN, yy + Inches(0.08), Inches(2.45), Inches(0.35), name, size=11.5, color=DEEPNAVY, bold=True)
    label(s, MARGIN, yy + Inches(0.42), Inches(2.45), Inches(0.4), f"Owner: {owner}", size=10, color=MIDGREY)
    for q0, nq, text in bars:
        col = [NAVY2, TAGBLUE, GREEN, DEEPGREEN][[0, 2, 6, 10].index(q0)]
        shp = rect(s, gx + qw * q0 + Inches(0.03), yy + Inches(0.12), qw * nq - Inches(0.09), lh - Inches(0.24), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        shape_text(shp, text, size=9.5, color=WHITE, bold=False, align=PP_ALIGN.LEFT, margins=(0.08, 0.02, 0.06, 0.02))
    # conditional P4 marker
    rect(s, gx + qw * 16 + Inches(0.03), yy + Inches(0.12), qw - Inches(0.09), lh - Inches(0.24), RGBColor(0xE1, 0xE6, 0xED), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
# dependencies row
dy = ly2 + 5 * lh + Inches(0.05)
label(s, MARGIN, dy, SLIDE_W - 2 * MARGIN, Inches(0.5), "Dependencies: gate 1 needs the regulator view and a station partner letter; Phase 1 needs car-size cylinder samples (Luxfer, Quantum, NPROXX, Time Technoplast) and hydrogen injectors; Phase 2 needs the MNRE sanction, PESO approvals and corridor station timing; Phase 3 needs the fuel contract and an insurer endorsement. Phase 4 (from Oct 2030, grey) is conditional: 1,000 to 1,500 cars at up to 8 hubs.", size=10.5, color=DARKGREY)
# gate table and funding
gy = dy + Inches(0.58)
rows = [["Gate", "Proceed if", "Otherwise"], ["G1 Mar 2027", "Written route from regulator and test agency; station partner secured; demand evidence from 3 fleets", "Stop"], ["G2 Mar 2028", "At or below 2.2 kg/100 km; NOx at or below 60 mg/km on cycle; no abnormal combustion in 500 h; safety tests passed", "Redesign once, then stop"], ["G3 Mar 2029", "Kit approval granted; dispensing point commissioned within 6 months of plan", "Pause"], ["G4 Sep 2030", "Availability 95 percent; zero reportable hydrogen incidents; measured ₹/km published; delivered hydrogen at or below ₹250/kg with a credible path to ₹190/kg", "Stop and publish"]]
table2(s, MARGIN, gy, Inches(13.2), Inches(1.9), rows, col_widths=[1.0, 5.2, 1.3], font_size=10, header_size=10.5, row_heights=[Inches(0.28), Inches(0.38), Inches(0.4), Inches(0.36), Inches(0.44)])
fx2 = MARGIN + Inches(13.45); fw2 = SLIDE_W - MARGIN - fx2
cyf = section_title(s, fx2, gy, fw2, f"Proposed funding of ₹{tot_cr:.1f} crore (shares are proposals)")
fund = [("SRDI and Maruti Suzuki", 0.40, NAVY2), ("Public: MNRE pilot or R&D scheme, state policy", 0.35, GREEN), ("OMC partner", 0.15, TAGBLUE), ("Fleet partner", 0.10, AMBER)]
bx_ = fx2; bw_tot = fw2
for t, sh, col in fund:
    ww = int(bw_tot * sh); rect(s, bx_, cyf, ww, Inches(0.38), col)
    label(s, bx_, cyf, ww, Inches(0.38), f"{sh*100:.0f}%", size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE); bx_ += ww
label(s, fx2, cyf + Inches(0.45), fw2, Inches(1.2), " | ".join(f"{t} ₹{sh*tot_cr:.1f} cr" for t, sh, _ in fund) + ". Why Suzuki's role is plausible: CNG engine and calibration know-how, 70 percent of India's CNG cars, its own hydrogen combustion prototypes (Swift with AVL, Burgman). No endorsement is implied.", size=10, color=DARKGREY)
footer2(s, "Sources: MNRE scheme guidelines, PIB 14 Feb 2024 (buses, trucks and 4-wheelers) [PIB 2006052]; ARAI Phase-II call for proposals (Sep to Dec 2025); test facilities [W4-S14, S15]; CNG kit approval precedent [W2-S3, S8]; PESO rules [W2-S24]; Pilot_Budget sheet.", 8)

# =============================================================== SLIDE 9
s = blank_slide(prs)
y0 = header2(s, f"Impact is real but small by 2030: about {round(M['fleet2030'], -2):,.0f} cars and {EM['fleet_t']/1000:.1f} kt CO2e a year, conditional on the gates. The decision requested is ₹{ph_cr[0]:.1f} crore for Phase 0",
             "Matched service (same car, same duty) for every comparison; projections are labelled as modelled, pilot target or conditional scenario, never as achieved results.")
lx, lw = MARGIN, Inches(6.5)
cy = section_title(s, lx, y0, lw, "Emissions per km, well-to-wheel with manufacturing (model)")
bar_chart(s, lx, cy, lw, Inches(3.5), ["Continue on CNG", "Retrofit, green H2", "Retrofit, grey H2", "BEV, 2025 grid", "BEV, 2030 grid"], [("g CO2e per km", (round(EM['A']), round(EM['Bgreen']), round(EM['Bgrey']), round(EM['C25']), round(EM['C30'])))], [GREEN], number_format='0', font_size=10, legend=False, label_size=11, value_axis_title="g CO2e per km")
label(s, lx, cy + Inches(3.55), lw, Inches(0.95), f"Green-hydrogen retrofit is {EM['redB']*100:.0f} percent below CNG with petrol backup km included; grid-electrolysis hydrogen would be {EM['Bgrid']:.0f} g/km, four times CNG. Reversal: hydrogen above {EM['cithr']:.1f} kg CO2e/kg makes the retrofit worse than CNG; grid below {EM['gridthr']:.2f} kg/kWh at the meter makes the BEV cleaner (CEA path about 0.5 by 2032). NOx remains.", size=10.5, color=MIDGREY)
cy2 = section_title(s, lx, cy + Inches(4.6), lw, "Scenario table (per car unless stated)")
rows = [["Metric", "Baseline: keep CNG", "Modelled retrofit, base", "Pilot target 2029-30", "Conditional 2030 (about 1,400 cars)"],
        ["CO2e, g/km", f"{EM['A']:.0f}", f"{EM['Bgreen']:.0f}", "at or below 95", f"{EM['fleet_t']:,.0f} t/yr avoided"],
        ["Hydrogen use", "0", f"{IN['h2_cons']['active']:.2f} kg/100 km", "at or below 2.2", f"{EM['fleet_h2']:,.0f} t/yr demand"],
        ["Cost per km, all-in", f"₹{S['lcA']:.1f}", f"₹{S['lcB']:.1f}", "measured and published", "depends on hydrogen price gate"],
        ["Availability", "reference", "not yet measured", "at or above 95 percent", "at or above 95 percent"]]
table2(s, lx, cy2, lw, Inches(2.2), rows, col_widths=[1.3, 1.2, 1.35, 1.3, 1.6], font_size=10, header_size=10, row_heights=[Inches(0.5)] + [Inches(0.4)] * 4)
# middle: scorecard + risks
mx, mw = lx + lw + Inches(0.3), Inches(6.5)
cy = section_title(s, mx, y0, mw, "Pilot scorecard: metric, success threshold, measurement")
rows = [["Metric", "Threshold (gate 4)", "How measured"], ["Consumption", "at or below 2.2 kg/100 km in fleet use", "Dispenser mass logs, CAN flow"], ["Range per fill", "at or above 55 km", "Trip logs, GPS"], ["Hydrogen share of km", "at or above 60 percent", "Telematics, fuel mode log"], ["Availability", "at or above 95 percent", "Downtime by cause"], ["Safety", "zero reportable incidents; all alarms closed", "Incident register, sensor logs"], ["NOx", "at or below 60 mg/km on cycle", "Test agency, quarterly"], ["Maintenance cost", "within ₹1.2/km", "Job cards"], ["Refuelling time", "under 6 min per fill", "Dispenser logs"], ["Delivered hydrogen", "at or below ₹250/kg, path to ₹190", "Fuel invoices, station utilisation"], ["Acceptance", "fleet managers and drivers accept power and boot", "Quarterly interviews, willingness to pay"]]
table2(s, mx, cy, mw, Inches(4.0), rows, col_widths=[1.3, 2.1, 1.7], font_size=10, header_size=10.5, row_heights=[Inches(0.3)] + [Inches(0.36)] * 10)
cy2 = section_title(s, mx, cy + Inches(4.3), mw, "Major unresolved risks (ranked)")
risks = [("Delivered hydrogen price", "₹360 to ₹1,430/kg across sources; decides the scale case", "Price gate; take-or-pay; pilot fuel as programme cost"), ("Approval route", "No instrument for a hydrogen retrofit kit", "Gate 1 consultation; trial under trade certificate"), ("NOx on transients", "Lean window lost under acceleration", "NOx trap provision; load limiting; cycle test at gate 2"), ("Cylinder supply", "No car-size PESO-approved cylinder", "Samples from 4 suppliers in Phase 0; PESO prototype approval"), ("Customer demand", "Hypothesis; power and boot loss may be refused", "30 interviews; gate 1 needs 3 fleets")]
rows = [["Risk", "Why it matters", "Mitigation"]] + [list(r) for r in risks]
table2(s, mx, cy2, mw, Inches(2.2), rows, col_widths=[1.2, 2.0, 1.9], font_size=10, header_size=10.5, row_heights=[Inches(0.3)] + [Inches(0.38)] * 5)
# right: decision
rx, rw = mx + mw + Inches(0.3), SLIDE_W - MARGIN - (mx + mw + Inches(0.3))
cy = section_title(s, rx, y0, rw, "Decision requested")
dec = [(f"Approve Phase 0: ₹{ph_cr[0]:.1f} crore, six months", "regulator consultations, 30 fleet interviews, telematics on 20 cars, boot mock-up; no commitment beyond gate 1"),
       ("Endorse a scheme application", "MNRE transport pilot scheme, 4-wheeler category, with ARAI as implementing agency; Gujarat station subsidy for the second hub"),
       ("Nominate roles", "SRDI engineering lead; fleet and fuel partner shortlist (proposed IOCL Faridabad; PSU and airport-contract fleets)"),
       ("Start HCNG in parallel", "where a CGD company blends to IS 17314: legal since 2020, no vehicle change, about 6 percent of energy from hydrogen; an air-quality step")]
for i, (t, d_) in enumerate(dec):
    ry = cy + i * Inches(1.0)
    numbered_circle(s, rx, ry + Inches(0.04), Inches(0.36), i + 1, fill=GREEN, size=12)
    label(s, rx + Inches(0.48), ry, rw - Inches(0.5), Inches(0.34), t, size=11.5, color=DEEPNAVY, bold=True)
    label(s, rx + Inches(0.48), ry + Inches(0.32), rw - Inches(0.5), Inches(0.68), d_, size=10.5, color=DARKGREY)
cy2 = section_title(s, rx, cy + Inches(4.15), rw, "What would make us stop")
stops = ["No written regulator route by March 2027", "Backfire or NOx not controlled within 500 test-cell hours", "Delivered hydrogen above ₹250/kg in 2030 with no path to ₹190/kg", "Fleets refuse the power and boot loss in Phase 0 interviews"]
for i, t in enumerate(stops):
    ry = cy2 + i * Inches(0.42)
    rect(s, rx, ry + Inches(0.1), Inches(0.18), Inches(0.18), REDC)
    label(s, rx + Inches(0.3), ry, rw - Inches(0.3), Inches(0.42), t, size=11, color=DARKGREY, anchor=MSO_ANCHOR.MIDDLE)
callout(s, rx, cy2 + Inches(1.85), rw, Inches(1.35), "H2-SHIFT keeps India's hydrogen retrofit option open at low cost, states the price at which it works, and hands the 2030 scale decision to evidence.", fill=DARKGREEN, size=12, bold=True, color=WHITE)
footer2(s, "Sources: CEA CO2 baseline v21 and National Electricity Plan [W4-S59, S60]; IPCC fuel factors [W4-S62]; ICCT 2025 lifecycle and battery factors [W4-S63]; hydrogen carbon intensities [W4-S64, S65]; DOE NOx page [W2-S52]; MNRE scheme scope [PIB 2006052]; H-CNG notification [W2-S1]; Emissions, Market and Pilot_Budget sheets.", 9)

# =============================================================== SLIDE 10
s = blank_slide(prs)
y0 = header2(s, "Key sources by theme (IDs as used on the slides) and annexure index", "Research cut-off 14 September 2026. The full register (309 entries with URLs, dates, sections and limitations) is in Annexure A2 and the workbook Sources sheet.")
groups = [
 ("Policy, standards and approval", ["W2-S13 MoRTH GSR 746(E), 16 Oct 2023: Rule 125M, AIS 195:2023 (Amd 1, May 2025)", "W2-S11 MoRTH GSR 579(E), 23 Sep 2020: AIS 157:2020; scope excludes retrofitted FCEVs", "W2-S1 MoRTH GSR 585(E), 25 Sep 2020: H-CNG per IS 17314:2019", "W2-S3, S9 CMVR Rules 115 and 126 (CNG/LPG kit provisions, retrofit proviso)", "W2-S24, S27 Gas Cylinders (Amendment) Rules 2025, GSR 225(E); PESO Form-H licensing (May 2025)", "W2-S42 to S45 UN GTR 13 Amend. 1 (2023); H2Tools and DOE properties; NFPA 2 (2023)"]),
 ("Market, fleet and customer", ["W3-S2 PNGRB citing Vahan: 81.95 lakh CNG vehicles (Mar 2025)", "W3-S4, S6 MoPNG: 8,150 plus CNG stations; 17,700 to 18,000 by 2030 (targets)", "W3-S10, S12 Maruti Suzuki CNG sales: 6.12 to 6.2 lakh FY25, above 7 lakh FY26; 71.6 percent segment share", "W3-S17, S19 Delhi aggregator scheme 2023; CAQM NCR direction (Jan 2026)", "W4-S30 NITI Aayog (Feb 2026): 4W taxi mileage 70,000 km/yr", "W3-S27, S30 Tata Tiago EV, Punch EV, Xpres-T prices; Autocar India real-world tests"]),
 ("Vehicle, components and combustion", ["W1-S3a to S3c Maruti Suzuki Tour S and Dzire specifications and releases", "W1-S2 Luxfer G-Stor Pro H2 table (V060H35: 60 L, 350 bar, 36 kg, 1.45 kg)", "W1-S1 Hexagon Purus Type 4 cylinder table (Rev 08/22)", "W1-S9a, S9b, S9h PFI air displacement and power loss; lambda-NOx behaviour (IJHE 2024-25; Onorati et al.)", "W1-S11a ISO 11439 hydrogen limit for steel CNG cylinders", "W4-S1, S2 US DOE records 19008 and 24006: Type IV storage system cost"]),
 ("Hydrogen supply and cost", ["PIB 2107795 (3 Mar 2025): 5 pilots, 37 vehicles, 9 stations, 10 routes, ₹208 crore", "PIB 2006052 (14 Feb 2024): scheme guidelines, buses, trucks and 4-wheelers, ₹496 crore; W3-S37 PIB 29 Jul 2026: 12 projects, 70 vehicles, 16 stations", "W4-S22 CEEW 2025: delivered green hydrogen USD 4.2/kg in 2030; above USD 10/kg today", "W4-S33, S38, S57 CEEW, SIGHT ammonia auctions, IEA GHR: production cost evidence", "W3-S39 Numaligarh refinery tender ₹279/kg; W4-S46 ICCT station cost by utilisation", "W3-S46, W4-S81 IOCL R&D Faridabad hydrogen dispensing"]),
 ("Emissions", ["W4-S59, S60 CEA CO2 Baseline Database v21 (0.710 t/MWh, FY2024-25); National Electricity Plan 2022-32", "W4-S62 IPCC 2006 fuel factors (petrol, natural gas)", "W4-S63 ICCT 2025 lifecycle factors; battery about 70 kg CO2e/kWh", "W4-S64, S65 IEA hydrogen carbon intensities; MNRE Green Hydrogen Standard (2 kg CO2e/kg)", "W2-S52 US DOE: hydrogen combustion and NOx"]),
 ("Global programmes", ["W3-S51, S52 Toyota Mirai sales; Hyundai Nexo", "W3-S55, S56 Toyota hydrogen HiAce pilot; GR Corolla H2", "Fuel Cells Works, RushLane (May 2026): Suzuki-AVL hydrogen Swift prototype; W3-S58 Suzuki Burgman", "W3-S59 to S63 Cummins B6.7H India; JCB type approval; KEYOU; MAN; ULEMCo", "W3-S64 to S67 Stellantis (Jul 2025); Shell California (Feb 2024); H2 Mobility Germany (2025); GM-Honda"]),
]
gw_ = (Inches(13.1) - Inches(0.4)) / 3; gh_ = Inches(4.1)
for i, (t, items) in enumerate(groups):
    gxx = MARGIN + (i % 3) * (gw_ + Inches(0.2)); gyy = y0 + (i // 3) * (gh_ + Inches(0.2))
    cyg = section_title(s, gxx, gyy, gw_, t, size=12.5)
    textbox(s, gxx, cyg, gw_, gh_ - Inches(0.5), [(it, {'bullet': True}) for it in items], size=12, color=DARKGREY, line_spacing=1.04, space_after=4)
rx = MARGIN + Inches(13.3); rw = SLIDE_W - MARGIN - rx
cy = section_title(s, rx, y0, rw, "Annexure index (separate PDF and workbook)")
rows = [["Section", "Content"], ["A1", "Claim audit: Round 1 versus evidence, with corrections"], ["A2", "Assumptions, evidence classes, full source register"], ["A3", "Cost model: BOM, TCO cash flows, break-even, sensitivities"], ["A4", "Packaging, energy balance and performance calculations"], ["A5", "Regulatory compliance matrix and hazard analysis"], ["A6", "Market sizing method and funnel"], ["A7", "Global benchmark evidence"], ["A8", "Pilot design, instrumentation, first 90 days"], ["A9", "Phase budgets, responsibilities, funding"], ["A10", "Judge Q&A, validation log, change log"], ["Workbook", "H2SHIFT model: 11 formula-driven sheets, recalculated"]]
table2(s, rx, cy, rw, Inches(4.3), rows, col_widths=[0.9, 4.3], font_size=10.5, header_size=10.5, row_heights=[Inches(0.32)] * 12)
cls = {}
for k_, v_ in IN.items(): cls[v_["cls"]] = cls.get(v_["cls"], 0) + 1
cyc = section_title(s, rx, cy + Inches(4.5), rw, f"Model inputs by evidence class ({sum(cls.values())} inputs)")
bx_ = rx; tot_in = sum(cls.values())
for k_, col in (("A", GREEN), ("B", TAGBLUE), ("C", AMBER), ("D", GREYC)):
    ww = int(rw * cls.get(k_, 0) / tot_in)
    if ww > 0:
        rect(s, bx_, cyc, ww, Inches(0.36), col); label(s, bx_, cyc, ww, Inches(0.36), f"{k_} {cls.get(k_, 0)}", size=10.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE); bx_ += ww
label(s, rx, cyc + Inches(0.4), rw, Inches(0.3), " | ".join(f"{k_} {cls.get(k_, 0)}" for k_ in ("A", "B", "C", "D")) + " (A verified, B model, C assumption, D target)", size=10, color=MIDGREY)
callout(s, rx, cyc + Inches(0.75), rw, Inches(2.0), "Evidence classes: A verified external evidence; B model estimate; C explicit assumption or range; D proposed target. Many originals were inaccessible through the build environment's network policy; entries marked 'excerpt' or 'summary' rest on indexed page text and should be re-verified before external use. No customer interviews, letters of interest, test results or partnerships are claimed.", fill=YELLOW, size=10.5)
footer2(s, "Team Wookies, NIT Warangal. SRDI Igniters Innovation Challenge 2026, Engineering Track, Option 2: Hydrogen Retrofit Solution. Round 2 submission.", 10)
prs.save(OUT); print("saved", OUT, len(prs.slides))
