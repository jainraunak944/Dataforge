from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x00, 0x24, 0x44); GREEN = RGBColor(0x01, 0x6F, 0x49); GREY = RGBColor(0x3D, 0x3D, 0x3B)

def new_doc(landscape=True, base_size=10):
    d = Document()
    sec = d.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(1.6))
    sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.4)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(base_size)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, GREEN), ("Heading 3", 11, NAVY)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = color
        rf = h.element.rPr.rFonts
        for att in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(att)) is not None: del rf.attrib[qn(att)]
        rf.set(qn('w:ascii'), 'Arial'); rf.set(qn('w:hAnsi'), 'Arial'); rf.set(qn('w:eastAsia'), 'Arial'); rf.set(qn('w:cs'), 'Arial')
    return d

def para(d, text, size=None, bold=False, italic=False, color=None, align=None, space_after=4):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullets(d, items, size=None):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        if size: r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_); tcPr.append(shd)

def table(d, rows, widths=None, size=8.5, header=True, header_fill="002444", zebra=True, bold_first_col=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run("" if val is None else str(val)); r.font.size = Pt(size)
            if i == 0 and header:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, header_fill)
            elif bold_first_col and j == 0:
                r.bold = True
            if zebra and i > 0 and i % 2 == 0: shade(c, "F3F6F4")
            if widths: c.width = Cm(widths[j])
    if widths:
        for j, w in enumerate(widths):
            for cell in t.columns[j].cells: cell.width = Cm(w)
        grid = t._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn('w:gridCol'))):
            if j < len(widths): gc.set(qn('w:w'), str(int(widths[j] * 567)))
        tblPr = t._tbl.tblPr
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t
